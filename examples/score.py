"""Score a run against truth.csv: auto-load precision and what got flagged."""
import csv
import sys
from collections import defaultdict

import openpyxl

truth_path, review_path = sys.argv[1], sys.argv[2]
truth = {r["key"]: r for r in csv.DictReader(open(truth_path))}
wb = openpyxl.load_workbook(review_path, read_only=True)
lines = list(wb["BOQ Lines"].iter_rows(values_only=True))
h = lines[0]
idx = {k: h.index(k) for k in ("BOQ Key", "Status", "Value Loaded", "Work Type", "Reason", "Description")}
alloc = defaultdict(set)
for r in list(wb["Allocations"].iter_rows(values_only=True))[1:]:
    alloc[r[0]].add(r[4])
tot = right = wrong = flagged_ok = flagged_should_auto = missed_flags = 0
val_auto = val_all = 0.0
wrong_list, flag_list = [], []
for r in lines[1:]:
    key, st, val = r[idx["BOQ Key"]], r[idx["Status"]], r[idx["Value Loaded"]] or 0
    t = truth.get(key)
    if not t:
        continue
    tot += 1
    val_all += val
    exp = t["expected"]
    if st in ("AUTO", "RULE", "DECIDED"):
        val_auto += val
        if exp in ("PRELIM", "REVIEW"):
            missed_flags += 1
            wrong_list.append((key, r[idx["Description"]], "should be flagged", r[idx["Work Type"]]))
        elif set(exp.split(";")) == alloc[key]:
            right += 1
        else:
            wrong += 1
            wrong_list.append((key, r[idx["Description"]], f"exp {exp[:40]}", f"got {r[idx['Work Type']]} {sorted(alloc[key])[:3]}"))
    elif st == "REVIEW":
        if exp in ("PRELIM", "REVIEW"):
            flagged_ok += 1
        else:
            flagged_should_auto += 1
            flag_list.append((key, r[idx["Description"]], r[idx["Reason"]]))
print(f"lines scored {tot}  auto-correct {right}  auto-WRONG {wrong}  should-have-flagged {missed_flags}")
print(f"flagged (correctly) {flagged_ok}  flagged (could have auto) {flagged_should_auto}")
print(f"value auto-loaded {val_auto / val_all * 100:.1f}%")
for w in wrong_list[:40]:
    print("  WRONG", w)
for w in flag_list[:60]:
    print("  FLAG ", w)
