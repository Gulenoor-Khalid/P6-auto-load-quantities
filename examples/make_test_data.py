"""Generate a synthetic 17-floor tower: P6 XER + P6 Excel export + priced BOQ + ground truth."""
import csv
import random
import sys
from pathlib import Path

import openpyxl

random.seed(7)
OUT = Path(sys.argv[1] if len(sys.argv) > 1 else ".")
OUT.mkdir(parents=True, exist_ok=True)

FLOORS = ["B2", "B1", "GF"] + [f"L{n:02d}" for n in range(1, 18)] + ["RF"]
FIN_FLOORS = ["GF"] + [f"L{n:02d}" for n in range(1, 18)]


def flabel(f):  # how finishes activities name floors
    return {"GF": "Ground Floor", "B1": "Basement 1", "B2": "Basement 2", "RF": "Roof"}.get(f, f"Level {f[1:]}")


acts = []  # (id, name, wbs_path tuple, type, dur_hours)
wbs_ids = {}
seq = [1000]


def add(name, wbs, ttype="TT_Task", dur=40, code=None):
    seq[0] += 10
    aid = code or f"A{seq[0]}"
    acts.append((aid, name, wbs, ttype, dur))
    return aid


add("Notice to Proceed", ("General",), "TT_Mile", 0)
MOB = add("Mobilization & Site Setup", ("General",), dur=160)
PM = add("Project Management & Supervision", ("General",), "TT_LOE", 6000)
add("Structure Complete", ("General",), "TT_FinMile", 0)
add("Practical Completion", ("General",), "TT_FinMile", 0)
EXC = add("Bulk Excavation", ("Substructure", "Earthworks"), dur=480)
SHO = add("Shoring - Secant Pile Wall", ("Substructure", "Earthworks"), dur=400)
DEW = add("Dewatering", ("Substructure", "Earthworks"), dur=600)
PIL = {z: add(f"Bored Piling Zone {z}", ("Substructure", "Piling"), dur=320) for z in "AB"}
CAP = {z: add(f"Pile Caps Zone {z}", ("Substructure", "Foundations"), dur=160) for z in "AB"}
BLD = {z: add(f"Blinding Concrete - Zone {z}", ("Substructure", "Foundations"), dur=40) for z in "AB"}
RWP = add("Raft Waterproofing", ("Substructure", "Foundations"), dur=120)
RAF = {z: add(f"Raft Foundation - Zone {z}", ("Substructure", "Foundations"), dur=200) for z in "AB"}
BKF = add("Backfilling", ("Substructure", "Earthworks"), dur=160)

STR = {}  # (floor, element, zone) -> id
for f in FLOORS:
    for z in "AB":
        for t in ("Rebar", "Formwork", "Concrete"):
            STR[(f, f"Column {t}", z)] = add(f"{f} Zone {z} - Column {t}", ("Superstructure", f), dur=24)
    for t in ("Rebar", "Formwork", "Concrete"):
        STR[(f, f"Core Wall {t}", "")] = add(f"{f} - Core Wall {t}", ("Superstructure", f), dur=32)
    for z in "AB":
        for t in ("Formwork", "Rebar", "Concrete Pour"):
            STR[(f, f"Slab {t}", z)] = add(f"{f} Zone {z} - Slab {t}", ("Superstructure", f), dur=40)
    STR[(f, "Staircase Construction", "")] = add(f"{f} - Staircase Construction", ("Superstructure", f), dur=40)

FIN_TYPES = ["Blockwork", "Electrical First Fix", "Plumbing First Fix", "HVAC Ductwork & FCU Installation",
             "Fire Fighting Sprinkler Piping", "Fire Alarm & ELV Containment", "Internal Plaster", "Floor Screed",
             "Wet Area Waterproofing", "False Ceiling", "Floor & Wall Tiling", "Painting", "Doors & Ironmongery",
             "Kitchen & Wardrobe Joinery", "MEP Second Fix & Fixtures", "Final Cleaning & Snagging"]
FIN = {}
for f in FIN_FLOORS:
    for t in FIN_TYPES:
        FIN[(f, t)] = add(f"{t} - {flabel(f)}", ("Finishes", flabel(f)), dur=80)
BAS = {}
for f in ("B2", "B1"):
    for t in ("Basement Blockwork", "Basement MEP Services", "Epoxy Floor Coating", "Basement Painting"):
        BAS[(f, t)] = add(f"{t} - {f}", ("Finishes", flabel(f)), dur=80)
CW = {f: add(f"Curtain Wall Installation - {flabel(f)}", ("Facade",), dur=60) for f in FIN_FLOORS}
ROOF_WP = add("Roof Waterproofing & Insulation", ("Roof",), dur=80)
ROOF_TL = add("Roof Screed & Tiling", ("Roof",), dur=60)
ROOF_MEP = add("Rooftop MEP Plant Installation", ("Roof",), dur=200)
LIFTS = [add(f"Lift Installation - Lift {n}", ("Vertical Transportation",), dur=400) for n in range(1, 5)]
TC = add("Testing & Commissioning - MEP Systems", ("Commissioning",), dur=400)
EXT_H = add("External Works - Hardscape & Roads", ("External Works",), dur=300)
EXT_L = add("External Works - Landscaping", ("External Works",), dur=200)
EXT_B = add("Boundary Wall & Gates", ("External Works",), dur=150)
print("activities", len(acts))

# ---------------- XER ----------------
wbs_rows = [("100", "", "TWR", "Tower Project", "Y")]
wid = [100]


def wbs_id(path):
    parent = "100"
    for i in range(len(path)):
        key = path[:i + 1]
        if key not in wbs_ids:
            wid[0] += 1
            wbs_ids[key] = str(wid[0])
            wbs_rows.append((str(wid[0]), parent, key[-1][:10].upper().replace(" ", ""), key[-1], "N"))
        parent = wbs_ids[key]
    return parent


task_rows = []
for i, (aid, name, wbs, tt, dur) in enumerate(acts):
    task_rows.append((str(5000 + i), "1", wbs_id(wbs), aid, name, tt, str(dur), "TK_NotStart"))
xer = ["ERMHDR\t22.12\t2026-10-01\tProject\tadmin\tadmin\tdbx\tProject Management\tUSD"]
xer += ["%T\tPROJECT", "%F\tproj_id\tproj_short_name", "%R\t1\tTWR-17"]
xer += ["%T\tPROJWBS", "%F\twbs_id\tproj_id\tparent_wbs_id\twbs_short_name\twbs_name\tproj_node_flag"]
xer += [f"%R\t{w}\t1\t{p}\t{s}\t{n}\t{f}" for w, p, s, n, f in wbs_rows]
xer += ["%T\tTASK", "%F\ttask_id\tproj_id\twbs_id\ttask_code\ttask_name\ttask_type\ttarget_drtn_hr_cnt\tstatus_code"]
xer += ["%R\t" + "\t".join(r) for r in task_rows]
xer += ["%E"]
(OUT / "TWR17_schedule.xer").write_text("\n".join(xer), encoding="cp1252")

# P6 Excel export flavour
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "TASK"
ws.append(["task_code", "status_code", "wbs_id", "task_name", "task_type", "target_drtn_hr_cnt"])
ws.append(["Activity ID", "Activity Status", "WBS Code", "Activity Name", "Activity Type", "Original Duration(d)"])
TT = {"TT_Task": "Task Dependent", "TT_Mile": "Start Milestone", "TT_FinMile": "Finish Milestone", "TT_LOE": "Level of Effort"}
for aid, name, wbs, tt, dur in acts:
    ws.append([aid, "Not Started", "TWR." + ".".join(w[:10].upper().replace(" ", "") for w in wbs), name, TT[tt], dur / 8])
wb.save(OUT / "TWR17_P6_export.xlsx")

# ---------------- BOQ ----------------
truth = []  # key, expected activity ids (sorted ';'), note
wb = openpyxl.Workbook()
wb.remove(wb.active)
summ = wb.create_sheet("Summary")
summ.append(["BILL", "DESCRIPTION", "AMOUNT (USD)"])


class Bill:
    def __init__(self, title):
        self.ws = wb.create_sheet(title[:31])
        self.ws.append([f"PROJECT: 17-STOREY RESIDENTIAL TOWER - {title.upper()}"])
        self.ws.append([])
        self.ws.append(["Item", "Description", "Unit", "Qty", "Rate (USD)", "Amount (USD)"])
        self.n = 0
        self.total = 0.0
        self.title = title

    def head(self, text):
        self.ws.append([None, text])

    def line(self, desc, unit, qty, rate, expect, amount=None, item=None):
        self.n += 1
        amt = round(qty * rate, 2) if amount is None else amount
        self.ws.append([item or f"{chr(65 + (self.n - 1) % 26)}", desc, unit, qty, rate, amt])
        key = f"{self.ws.title}!R{self.ws.max_row}"
        truth.append((key, ";".join(sorted(expect)) if isinstance(expect, (list, set, tuple)) else expect, desc))
        self.total += amt
        return key

    def close(self, wrong_by=0.0):
        self.ws.append([None, "Total carried to Summary", None, None, None, round(self.total + wrong_by, 2)])
        summ.append([self.title, self.title, round(self.total, 2)])
        for col, w in zip("ABCDEF", (8, 70, 8, 12, 14, 18)):
            self.ws.column_dimensions[col].width = w


R = random.uniform
# Bill 1 prelims
b = Bill("Bill 1 - Preliminaries")
b.head("GENERAL REQUIREMENTS")
for d, amt in [("Site establishment, mobilisation and temporary facilities", 850000), ("Project manager and site staff for 30 months", 2400000),
               ("Contractor's all risk insurance", 310000), ("Performance bond", 280000), ("Temporary site power and water", 420000),
               ("Tower crane hire and operation, 2 nr", 1350000), ("Site security and hoarding", 260000), ("Health and safety provisions", 190000),
               ("Demobilisation and final clean", 120000)]:
    b.line(d, "Item", 1, amt, "PRELIM" if "mobilisation" not in d.lower() else MOB)
b.close()

# Bill 2 substructure
b = Bill("Bill 2 - Substructure")
b.head("EARTHWORKS")
b.line("Excavate to reduced levels in any material, max depth 12m", "m3", 38500, 9.5, [EXC])
b.line("Disposal of excavated material off site to approved tip", "m3", 38500, 6.2, [EXC])
b.line("Secant pile shoring wall 900mm dia, 18m deep", "m", 420, 2150, [SHO])
b.line("Dewatering system installation and operation", "Item", 1, 385000, [DEW])
b.line("Backfilling around basement walls with selected material, compacted", "m3", 4200, 14, [BKF])
b.head("PILING")
for z in "AB":
    b.line(f"Bored cast in situ piles 1000mm dia, Zone {z}", "m", 2600, 310, [PIL[z]])
b.head("FOUNDATIONS")
for z in "AB":
    b.line(f"Blinding concrete 75mm thick under raft, Zone {z}", "m2", 1400, 18, [BLD[z]])
    b.line(f"Pile cap concrete C45 including trimming pile heads, Zone {z}", "m3", 650, 245, [CAP[z]])
    b.line(f"Raft foundation concrete grade C45, Zone {z}", "m3", 3200, 235, [RAF[z]])
    b.line(f"High yield reinforcement to raft foundation, Zone {z}", "t", 420, 1180, [RAF[z]])
b.line("Bituminous waterproofing membrane to underside of raft", "m2", 2900, 32, [RWP])
b.close()

# Bill 3 superstructure - sections per floor
b = Bill("Bill 3 - Superstructure")
FLN = {"B2": "BASEMENT 2", "B1": "BASEMENT 1", "GF": "GROUND FLOOR", "RF": "ROOF"}
for f in FLOORS:
    b.head(FLN.get(f, f"LEVEL {f[1:]}"))
    both = lambda e: [STR[(f, e, z)] for z in "AB"]
    b.line("Reinforced concrete grade C60 in columns", "m3", round(R(60, 90), 1), 265, both("Column Concrete"))
    b.line("High yield bar reinforcement to columns, BS 4449 B500B", "t", round(R(18, 30), 2), 1210, both("Column Rebar"))
    b.line("Formwork to vertical sides of columns", "m2", round(R(500, 700)), 42, both("Column Formwork"))
    b.line("Reinforced concrete grade C50 in core walls", "m3", round(R(80, 110), 1), 255, [STR[(f, "Core Wall Concrete", "")]])
    b.line("Reinforcement to core walls", "t", round(R(14, 22), 2), 1205, [STR[(f, "Core Wall Rebar", "")]])
    b.line("Formwork to faces of core walls including openings", "m2", round(R(800, 1100)), 45, [STR[(f, "Core Wall Formwork", "")]])
    for z in "AB":
        b.line(f"Concrete grade C40 in suspended slab, Zone {z}", "m3", round(R(180, 240), 1), 228, [STR[(f, "Slab Concrete Pour", z)]])
        b.line(f"Bar reinforcement to suspended slab, Zone {z}", "t", round(R(20, 30), 2), 1190, [STR[(f, "Slab Rebar", z)]])
        b.line(f"Soffit formwork to suspended slab, Zone {z}", "m2", round(R(700, 900)), 38, [STR[(f, "Slab Formwork", z)]])
    b.line("In-situ staircase flights and landings complete", "Nr", 2, 9800, [STR[(f, "Staircase Construction", "")]])
b.close(wrong_by=15000.0)  # planted subtotal mismatch

# Bill 4 masonry & finishes
b = Bill("Bill 4 - Masonry & Finishes")
b.head("MASONRY")
b.line("200mm hollow concrete blockwork to internal partitions, Levels 1 to 17", "m2", 52000, 28, [FIN[(f, "Blockwork")] for f in FIN_FLOORS if f != "GF"])
b.line("200mm hollow blockwork walls, ground floor", "m2", 3100, 28, [FIN[("GF", "Blockwork")]])
b.line("Blockwork to basement plant rooms, B1 and B2", "m2", 2400, 27, [BAS[(f, "Basement Blockwork")] for f in ("B1", "B2")])
b.head("PLASTERING")
b.line("Internal cement plaster 15mm to blockwork walls, Ground to Level 17", "m2", 118000, 9.5, [FIN[(f, "Internal Plaster")] for f in FIN_FLOORS])
for f in FIN_FLOORS:
    b.head(f"FINISHES - {FLN.get(f, 'LEVEL ' + f[1:])}")
    b.line("Sand cement screed 75mm to floors", "m2", round(R(1100, 1300)), 14, [FIN[(f, "Floor Screed")]])
    b.line("Porcelain floor tiles 600x600mm to apartments", "m2", round(R(900, 1100)), 46, [FIN[(f, "Floor & Wall Tiling")]])
    b.line("Ceramic wall tiles to bathrooms and kitchens", "m2", round(R(600, 800)), 38, [FIN[(f, "Floor & Wall Tiling")]])
    b.line("Porcelain skirting 100mm high", "m", round(R(1400, 1700)), 7.5, [FIN[(f, "Floor & Wall Tiling")]])
    b.line("Liquid applied waterproofing membrane to wet areas", "m2", round(R(250, 320)), 24, [FIN[(f, "Wet Area Waterproofing")]])
    b.line("Suspended gypsum board false ceiling with access panels", "m2", round(R(900, 1100)), 33, [FIN[(f, "False Ceiling")]])
    b.line("Emulsion paint to plastered walls, 3 coats", "m2", round(R(5500, 6500)), 4.2, [FIN[(f, "Painting")]])
    b.line("Emulsion paint to ceilings, 2 coats", "m2", round(R(900, 1100)), 3.9, [FIN[(f, "Painting")]])
for f in ("B1", "B2"):
    b.line(f"Epoxy floor coating to car park, Basement {f[1]}", "m2", 5200, 19, [BAS[(f, "Epoxy Floor Coating")]])
    b.line(f"Emulsion paint to basement walls and soffits, Basement {f[1]}", "m2", 7400, 3.1, [BAS[(f, "Basement Painting")]])
b.head("ROOF")
b.line("Roof waterproofing membrane with 50mm XPS insulation", "m2", 1250, 58, [ROOF_WP])
b.line("Roof screed to falls and porcelain roof tiles", "m2", 1250, 62, [ROOF_TL])
b.close()

# Bill 5 doors & joinery
b = Bill("Bill 5 - Doors & Joinery")
for f in FIN_FLOORS:
    b.head(f"{FLN.get(f, 'LEVEL ' + f[1:])}")
    b.line("Type D1 solid core timber flush door 900x2100 with ironmongery", "Nr", random.randint(28, 36), 640, [FIN[(f, "Doors & Ironmongery")]])
    b.line("Type FD2 60 minute fire rated door with frame", "Nr", random.randint(6, 10), 1150, [FIN[(f, "Doors & Ironmongery")]])
    b.line("Kitchen base and wall cabinets with quartz worktop", "Nr", 8, 5200, [FIN[(f, "Kitchen & Wardrobe Joinery")]])
    b.line("Built-in wardrobes, laminated finish", "Nr", 16, 1350, [FIN[(f, "Kitchen & Wardrobe Joinery")]])
b.close()

# Bill 6 facade
b = Bill("Bill 6 - Facade")
for f in FIN_FLOORS:
    b.line(f"Unitised aluminium curtain wall system, {flabel(f)}", "m2", round(R(520, 600)), 610, [CW[f]])
b.line("Aluminium louvres to rooftop plant screen", "m2", 300, 280, "REVIEW")
b.close()

# Bill 7 MEP
b = Bill("Bill 7 - MEP Services")
for f in FIN_FLOORS:
    b.head(f"{FLN.get(f, 'LEVEL ' + f[1:])}")
    E = lambda t: [FIN[(f, t)]]
    b.line("PVC conduits and wiring for lighting and power points", "Item", 1, round(R(60000, 70000), 2), E("Electrical First Fix"))
    b.line("Distribution boards and sub-main cables", "Item", 1, round(R(25000, 30000), 2), E("Electrical First Fix"))
    b.line("LED light fittings type L1 including installation", "Nr", random.randint(180, 220), 85, E("MEP Second Fix & Fixtures"))
    b.line("uPVC drainage pipework and fittings", "m", round(R(700, 800)), 34, E("Plumbing First Fix"))
    b.line("PPR water supply piping with insulation", "m", round(R(900, 1000)), 22, E("Plumbing First Fix"))
    b.line("Sanitary ware: WC, wash basin and shower mixer sets", "Set", 16, 1450, E("MEP Second Fix & Fixtures"))
    b.line("Galvanised sheet metal ductwork with thermal insulation", "m2", round(R(800, 900)), 48, E("HVAC Ductwork & FCU Installation"))
    b.line("Fan coil units, ceiling concealed, 2-row", "Nr", random.randint(30, 36), 1150, E("HVAC Ductwork & FCU Installation"))
    b.line("Sprinkler heads and branch piping", "Nr", random.randint(140, 170), 95, E("Fire Fighting Sprinkler Piping"))
    b.line("Fire alarm smoke detectors and loop cabling", "Nr", random.randint(60, 75), 120, E("Fire Alarm & ELV Containment"))
    b.line("Data and TV outlets with structured cabling", "Nr", random.randint(60, 75), 95, E("Fire Alarm & ELV Containment"))
b.head("CENTRAL PLANT")
b.line("Air cooled chillers 450TR, 3 nr on roof", "Nr", 3, 285000, [ROOF_MEP])
b.line("Booster and fire pump sets", "Set", 2, 96000, [ROOF_MEP])
b.line("Passenger lifts 21 stops, 1.6 m/s", "Nr", 4, 178000, LIFTS)
b.line("Testing and commissioning of all MEP systems", "Item", 1, 210000, [TC])
b.line("Basement MEP services including car park ventilation", "Item", 1, 640000, [BAS[(f, "Basement MEP Services")] for f in ("B1", "B2")])
b.line("Sundry builder's work in connection with services", "Item", 1, 145000, "REVIEW")
b.close()

# Bill 8 external
b = Bill("Bill 8 - External Works")
b.line("Asphalt paving to access roads and car parking", "m2", 6200, 38, [EXT_H])
b.line("Interlocking concrete pavers to walkways", "m2", 2100, 42, [EXT_H])
b.line("Soft landscaping, planting and irrigation", "m2", 3800, 55, [EXT_L])
b.line("Boundary wall 2.4m high with steel gates", "m", 420, 680, [EXT_B])
b.line("Signage and wayfinding", "Item", 1, 65000, "REVIEW")
b.close()

# Bill 9 provisional sums
b = Bill("Bill 9 - Provisional Sums")
for d, a in [("Provisional sum for utility connections", 450000), ("Provisional sum for artwork", 80000),
             ("Contingency", 900000)]:
    b.line(d, "Sum", 1, a, "PRELIM")
b.close()

# planted BOQ errors
ws = wb["Bill 3 - Superstructure"]
ws["F12"] = ws["F12"].value + 1000          # arithmetic error
ws = wb["Bill 4 - Masonry & Finishes"]
ws["F20"] = round(ws["F20"].value * 1.1, 2)  # arithmetic error
ws = wb["Bill 7 - MEP Services"]
ws["F30"] = None                             # amount missing -> computed
ws = wb["Bill 5 - Doors & Joinery"]
ws.append(["X", "Extra over for omission of door type D1 at Level 10", "Nr", -4, 640, -2560])
truth.append((f"{ws.title}!R{ws.max_row}", "REVIEW", "negative"))
ws.append(["Y", "Door closers", "Nr", 120, None, None])  # unpriced
ws.append(["Z", "Door stops", "Nr", 300, "Incl", "Incl"])
ws = wb["Bill 6 - Facade"]
ws.append(["W", "Unitised aluminium curtain wall system, Level 18", "m2", 500, 610, 305000])
truth.append((f"{ws.title}!R{ws.max_row}", "REVIEW", "level 18 not in schedule"))
wb.save(OUT / "TWR17_BOQ.xlsx")

with open(OUT / "truth.csv", "w", newline="") as f:
    w = csv.writer(f)
    w.writerow(["key", "expected", "description"])
    w.writerows(truth)
print("boq lines", len(truth))
