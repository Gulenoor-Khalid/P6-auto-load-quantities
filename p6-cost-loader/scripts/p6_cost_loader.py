#!/usr/bin/env python3
"""P6 Cost Loader: cost-load a Primavera P6 schedule from a priced BOQ.

Matches every priced BOQ line to a P6 work type, splits its value across the
floor/zone activities of that work type to the exact cent, flags uncertain
matches and BOQ arithmetic errors for the planner, reconciles every cent, and
writes a P6 resource-assignment import file plus a review workbook.

Commands
  inspect   Show how the BOQ and schedule were parsed (columns, floors, work types).
  run       Match, split, reconcile and write the outputs.

Examples
  python p6_cost_loader.py inspect --boq BOQ.xlsx --p6 schedule.xer
  python p6_cost_loader.py run --boq BOQ.xlsx --p6 schedule.xer --out out/
  python p6_cost_loader.py run ... --decisions out/cost_loading_review.xlsx

Only needs Python 3.8+ and openpyxl.
"""
import argparse
import csv
import fnmatch
import json
import math
import re
import sys
import time
from collections import Counter, defaultdict
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import column_index_from_string, get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required: pip install openpyxl")

CENT = Decimal("0.01")
ZERO = Decimal("0")
HERE = Path(__file__).resolve().parent
LEXICON_PATH = HERE.parent / "assets" / "trade_lexicon.json"

LOADED = ("AUTO", "RULE", "DECIDED")


# --------------------------------------------------------------------------
# Small helpers
# --------------------------------------------------------------------------

def norm_space(s):
    return re.sub(r"\s+", " ", s).strip()


def cell_str(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def to_dec(v):
    """Parse a cell into Decimal, or None if it is not a number."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, int):
        return Decimal(v)
    if isinstance(v, float):
        if math.isnan(v) or math.isinf(v):
            return None
        return Decimal(repr(v))
    if isinstance(v, Decimal):
        return v
    s = str(v).strip().replace(",", "").replace(" ", "").replace(" ", "")
    s = re.sub(r"^[A-Za-z$€£¥₹]+\.?", "", s)  # currency prefix like AED, $, Rs.
    m = re.fullmatch(r"(\()?(-?\d*\.?\d+)(\))?", s)
    if not m:
        return None
    d = Decimal(m.group(2))
    return -d if m.group(1) and m.group(3) else d


def cents(d):
    return d.quantize(CENT, rounding=ROUND_HALF_UP)


def money(d):
    return f"{d:,.2f}"


def allocate(amount, weights):
    """Split amount across weights so the parts sum to amount exactly (largest remainder)."""
    total_c = int((amount / CENT).to_integral_value(ROUND_HALF_UP))
    sign = -1 if total_c < 0 else 1
    total_c = abs(total_c)
    weights = [w if w > 0 else ZERO for w in weights]
    wsum = sum(weights)
    if wsum <= 0:
        weights = [Decimal(1)] * len(weights)
        wsum = Decimal(len(weights))
    raw = [Decimal(total_c) * w / wsum for w in weights]
    base = [int(r) for r in raw]
    rest = total_c - sum(base)
    order = sorted(range(len(raw)), key=lambda i: (-(raw[i] - base[i]), i))
    for i in order[:rest]:
        base[i] += 1
    return [Decimal(sign * b) * CENT for b in base]


def read_tables(path):
    """Return {sheet_name: [row tuples]} for xlsx/xlsm/csv/tsv/txt."""
    p = Path(path)
    ext = p.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        out = {ws.title: [tuple(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}
        wb.close()
        return out
    if ext in (".csv", ".tsv", ".txt"):
        raw = p.read_bytes()
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        try:
            dialect = csv.Sniffer().sniff(text[:5000], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel_tab if ext == ".tsv" else csv.excel
        return {p.stem: [tuple(r) for r in csv.reader(text.splitlines(), dialect)]}
    if ext == ".xls":
        sys.exit(f"{p.name}: legacy .xls is not supported - open it in Excel and save as .xlsx")
    sys.exit(f"{p.name}: unsupported file type {ext}")


def norm_header(v):
    s = cell_str(v).lower()
    s = re.sub(r"\(.*?\)|\[.*?\]", " ", s)
    s = re.sub(r"[^a-z0-9/&_ ]+", " ", s)
    return norm_space(s)


def match_role(h, aliases):
    if not h:
        return None
    for role, al in aliases.items():
        if h in al:
            return role
    for role, al in aliases.items():
        if any(h.startswith(a + " ") for a in al if len(a) > 2):
            return role
    return None


def find_header(rows, aliases, required, last_wins=("amount",), scan=40):
    best = None
    for i, row in enumerate(rows[:scan]):
        roles = {}
        for j, v in enumerate(row):
            r = match_role(norm_header(v), aliases)
            if r and (r not in roles or r in last_wins):
                roles[r] = j
        if required(roles) and (best is None or len(roles) > len(best[1])):
            best = (i, roles)
    return best


def col_ref_to_index(ref, header_row):
    ref = ref.strip()
    if re.fullmatch(r"[A-Za-z]{1,3}", ref):
        return column_index_from_string(ref.upper()) - 1
    if ref.isdigit():
        return int(ref) - 1
    target = norm_header(ref)
    for j, v in enumerate(header_row):
        if norm_header(v) == target:
            return j
    sys.exit(f"Column '{ref}' not found in header row")


def get(row, idx):
    if idx is None or idx >= len(row):
        return None
    return row[idx]


# --------------------------------------------------------------------------
# Location parsing (floors / zones)
# --------------------------------------------------------------------------

FLOOR_ORDER_PREFIX = {"B": 0, "LG": 1, "GF": 2, "UG": 3, "MZ": 4, "P": 5, "L": 6, "RF": 9}


def floor_sort_key(f):
    if f.startswith("B") and f[1:].isdigit():
        return (0, -int(f[1:]))
    if f.startswith("P") and f[1:].isdigit():
        return (5, int(f[1:]))
    if f.startswith("L") and f[1:].isdigit():
        return (6, int(f[1:]))
    return (FLOOR_ORDER_PREFIX.get(f, 8), 0)


def L(n):
    return "GF" if int(n) == 0 else f"L{int(n):02d}"


ORD = r"(?:st|nd|rd|th)"
NOT_DIM = r"(?!\s*(?:mm|cm|m\b|m2|m3|x|×|\*|\.\d|mpa|kg|nr|no\b|thk|thick|dia))"  # "floor 600x600" is a size
RANGE_SEP = r"\s*(?:to|-|–|—|thru|through|~)\s*"
FLOOR_WORD = r"(?:levels?|lvls?|lev|floors?|flrs?|fl)"

FLOOR_RANGES = [
    (re.compile(rf"\b{FLOOR_WORD}\s*[-_.#]?\s*(\d{{1,3}}){RANGE_SEP}(?:{FLOOR_WORD}\s*|L-?)?(\d{{1,3}})\b", re.I), "num"),
    (re.compile(rf"\bL-?(\d{{1,3}}){RANGE_SEP}L-?(\d{{1,3}})\b", re.I), "num"),
    (re.compile(rf"\b(\d{{1,3}}){ORD}{RANGE_SEP}(\d{{1,3}}){ORD}\s*(?:floors?|levels?)\b", re.I), "num"),
    (re.compile(rf"\b(?:basements?|B)\s*-?(\d){RANGE_SEP}(?:basement|B)\s*-?(\d)\b", re.I), "bsmt"),
    (re.compile(rf"\b(GF|ground\s+(?:floor|level)|ground){RANGE_SEP}(?:{FLOOR_WORD}\s*|L-?)?(\d{{1,3}}|roof|RF)\b", re.I), "gf"),
]
FLOOR_SINGLE_STRONG = [
    (re.compile(r"\b(?:basement|bsmt)\s*[-_]?\s*(\d{1,2})\b", re.I), lambda m: f"B{int(m.group(1))}"),
    (re.compile(r"\b(?:lower\s+ground(?:\s+floor)?|LGF)\b", re.I), lambda m: "LG"),
    (re.compile(r"\b(?:upper\s+ground(?:\s+floor)?|UGF)\b", re.I), lambda m: "UG"),
    (re.compile(r"\b(?:ground\s+(?:floor|level)|G/F|G\.F\.?)", re.I), lambda m: "GF"),
    (re.compile(r"\bGF\b"), lambda m: "GF"),
    (re.compile(r"\b(?:mezzanine|mezz)\b", re.I), lambda m: "MZ"),
    (re.compile(r"\bpodium\s*[-_]?\s*(\d{1,2})\b", re.I), lambda m: f"P{int(m.group(1)):02d}"),
    (re.compile(rf"\b(\d{{1,3}}){ORD}\s*(?:floor|flr|level)\b", re.I), lambda m: L(m.group(1))),
    (re.compile(rf"\b{FLOOR_WORD}\s*[-_.#]?\s*(\d{{1,3}})\b{NOT_DIM}", re.I),
     lambda m: L(m.group(1)) if int(m.group(1)) <= 120 else None),
    (re.compile(r"\b(?:roof(?:\s*top)?|rooftop|roof\s+level)\b", re.I), lambda m: "RF"),
    (re.compile(r"\bRF\b"), lambda m: "RF"),
]
# Bare codes like L05, B2, LG, MZ: trusted in schedules; in BOQ text only when the
# floor exists in the schedule and the code is not a type/grade mark ("Type B2 door").
FLOOR_SINGLE_WEAK = [
    (re.compile(r"\bL-?(\d{1,3})\b"), lambda m: L(m.group(1))),
    (re.compile(r"\bB-?(\d)\b"), lambda m: f"B{int(m.group(1))}"),
    (re.compile(r"\bP-?(\d{1,2})\b"), lambda m: f"P{int(m.group(1)):02d}"),
    (re.compile(r"\b(LG|UG|MZ)\b"), lambda m: m.group(1)),
]
ZONE_RANGE = re.compile(r"\bzones?\s*([A-Z]\d{0,2}|\d{1,2})\s*(&|and|,|/|-|to)\s*([A-Z]\d{0,2}|\d{1,2})\b", re.I)
ZONE_STRONG = re.compile(r"\b(?:zone|sector|area|pour)\s*[-_#:]?\s*([A-Z]\d{0,2}|\d{1,2})\b", re.I)
ZONE_WEAK = re.compile(r"\bZ-?(\d{1,2}|[A-H])\b")
MARK_WORDS = {"type", "grade", "ref", "mark", "class", "no", "item", "drawing", "dwg", "detail", "sheet", "section", "bar", "size", "fy", "c"}


def _prev_word(text, pos):
    m = re.search(r"([A-Za-z]+)\W*$", text[:pos])
    return m.group(1).lower() if m else ""


def parse_location(text, universe=None, trusted=False):
    """Return (floors, zones, spans, unknown_floors) found in text.

    trusted=True for schedule names (accept bare codes); False for BOQ text, where bare
    codes need to exist in `universe` (the schedule's floors) to count.
    """
    floors, zones, spans, unknown = set(), set(), [], set()
    taken = [False] * (len(text) + 1)

    def free(m):
        return not any(taken[m.start():m.end()])

    def take(m):
        for i in range(m.start(), m.end()):
            taken[i] = True
        spans.append((m.start(), m.end()))

    order = sorted(universe or [], key=floor_sort_key)

    def expand(a, b):
        if order and a in order and b in order:
            ia, ib = order.index(a), order.index(b)
            lo, hi = min(ia, ib), max(ia, ib)
            return set(order[lo:hi + 1])
        return {a, b}

    for rx, kind in FLOOR_RANGES:
        for m in rx.finditer(text):
            if not free(m):
                continue
            if kind == "num":
                a, b = int(m.group(1)), int(m.group(2))
                if a > b or b - a > 150:
                    continue
                got = {L(n) for n in range(a, b + 1)}
            elif kind == "bsmt":
                a, b = sorted((int(m.group(1)), int(m.group(2))))
                got = {f"B{n}" for n in range(a, b + 1)}
            else:
                end = m.group(2).upper()
                endf = "RF" if end in ("ROOF", "RF") else L(end)
                if order:
                    got = expand("GF", endf)
                elif end.isdigit():
                    got = {"GF"} | {L(n) for n in range(1, int(end) + 1)}
                else:
                    got = {"GF", endf}
            floors |= got
            take(m)

    def add_floor(f, strong):
        if universe is not None and f not in universe:
            if strong:
                unknown.add(f)
            return False
        floors.add(f)
        return True

    for rx, fn in FLOOR_SINGLE_STRONG:
        for m in rx.finditer(text):
            if free(m):
                f = fn(m)
                if f:
                    add_floor(f, True)
                    take(m)
    for rx, fn in FLOOR_SINGLE_WEAK:
        for m in rx.finditer(text):
            if not free(m):
                continue
            if not trusted and _prev_word(text, m.start()) in MARK_WORDS:
                continue
            f = fn(m)
            if trusted:
                floors.add(f)
                take(m)
            elif universe and f in universe:
                floors.add(f)
                take(m)

    for m in ZONE_RANGE.finditer(text):
        if not free(m):
            continue
        a, sep, b = m.group(1).upper(), m.group(2).lower(), m.group(3).upper()
        if sep in ("-", "to") and a.isdigit() and b.isdigit() and int(a) < int(b):
            zones |= {str(n) for n in range(int(a), int(b) + 1)}
        elif sep in ("-", "to") and len(a) == len(b) == 1 and a.isalpha() and b.isalpha() and a < b:
            zones |= {chr(c) for c in range(ord(a), ord(b) + 1)}
        else:
            zones |= {a.lstrip("0") or "0", b.lstrip("0") or "0"}
        take(m)
    for m in ZONE_STRONG.finditer(text):
        if free(m):
            zones.add(m.group(1).upper().lstrip("0") or "0")
            take(m)
    for m in ZONE_WEAK.finditer(text):
        if free(m) and (trusted or _prev_word(text, m.start()) not in MARK_WORDS):
            zones.add(m.group(1).upper().lstrip("0") or "0")
            take(m)
    return floors, zones, spans, unknown


def strip_spans(text, spans):
    out, last = [], 0
    for a, b in sorted(spans):
        if a < last:
            continue
        out.append(text[last:a])
        out.append(" ")
        last = b
    out.append(text[last:])
    return "".join(out)


def summarize_floors(floors):
    if not floors:
        return ""
    fl = sorted(floors, key=floor_sort_key)
    if len(fl) <= 3:
        return ", ".join(fl)
    return f"{fl[0]}..{fl[-1]} ({len(fl)})"


# --------------------------------------------------------------------------
# Text normalisation and scoring
# --------------------------------------------------------------------------

STOPWORDS = set("""
a an the and or of to in on at by for from with without into onto as per all any each every be is are was were this that these those
supply supplying install installing installation installed provide providing provision fix fixing fixed erect erecting construct constructing construction
including include includes incl inclusive complete completed etc approved specified specification spec specs drawing drawings dwg dwgs
work works item items nr no nos number qty quantity mm cm m lm rm m1 m2 m3 sqm sq cum cu kg kgs ton tons tonne tonnes t ls sum lump each ea
thick thk thickness wide width deep depth high height size sizes length type types grade class mpa n x
level levels floor floors zone zones area areas sector pour lvl typical general misc miscellaneous other others
""".split())
CANON_WEIGHT = 1.5


class Lexicon:
    def __init__(self, path):
        data = json.loads(Path(path).read_text(encoding="utf-8")) if Path(path).exists() else {}
        pairs = []
        for canon, phrases in data.items():
            if canon.startswith("_"):
                continue
            for p in set(phrases + [canon]):
                pairs.append((p.lower(), "_" + canon))
        self.rx = None
        self.map = defaultdict(list)  # a phrase may carry several terms: "sanitary ware" -> plumbing + secondfix
        for p, c in pairs:
            if c not in self.map[p]:
                self.map[p].append(c)
        if pairs:
            alt = "|".join(re.escape(p) for p in sorted(self.map, key=len, reverse=True))
            self.rx = re.compile(rf"(?<![a-z0-9])(?:{alt})(?![a-z0-9])")

    def canon(self, text):
        if not self.rx:
            return text
        return self.rx.sub(lambda m: " " + " ".join(self.map[m.group(0)]) + " ", text)


def stem(w):
    if len(w) > 4 and w.endswith("ies"):
        return w[:-3] + "y"
    if len(w) > 4 and (w.endswith("ches") or w.endswith("shes") or w.endswith("xes") or w.endswith("sses")):
        return w[:-2]
    if len(w) > 3 and w.endswith("s") and not w.endswith("ss"):
        return w[:-1]
    if len(w) > 6 and w.endswith("ing"):
        return w[:-3]
    return w


def tokens(text, lex):
    t = lex.canon(text.lower())
    out = []
    for w in re.findall(r"_?[a-z][a-z0-9_]*", t):
        if w.startswith("_"):
            out.append(w)
            continue
        if w in STOPWORDS or len(w) < 2:
            continue
        w = stem(w)
        if w not in STOPWORDS:
            out.append(w)
    return out


def add_tokens(vec, toks, weight, decay=0.0):
    """decay > 0 down-weights later words: BOQ descriptions lead with the item
    ("Fan coil units, ceiling concealed") and trail off into qualifiers."""
    for i, t in enumerate(toks):
        w = weight * max(0.6, 1.0 - decay * i) if decay else weight
        vec[t] += w * (CANON_WEIGHT if t.startswith("_") else 1.0)


# --------------------------------------------------------------------------
# Schedule (P6) reading
# --------------------------------------------------------------------------

class Activity:
    __slots__ = ("id", "name", "wbs_path", "wbs_code", "duration", "atype", "floors", "zones",
                 "worktype", "wbs_nonloc", "group", "kind", "status")

    def __init__(self, aid, name, wbs_path, wbs_code, duration, atype, status=""):
        self.id, self.name, self.wbs_path, self.wbs_code = aid, name, wbs_path, wbs_code
        self.duration, self.atype, self.status = duration, atype, status
        t = (atype or "").lower()
        if "mile" in t:
            self.kind = "MILESTONE"
        elif "loe" in t or "level of effort" in t:
            self.kind = "LOE"
        elif t in ("tt_wbs", "wbs summary"):
            self.kind = "WBS_SUMMARY"
        else:
            self.kind = "TASK"
        self.floors, self.zones = set(), set()
        self.worktype = self.wbs_nonloc = self.group = ""


def read_xer(path, project=None):
    raw = Path(path).read_bytes()
    for enc in ("utf-8", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    tables, cur, fields = {}, None, None
    for line in text.splitlines():
        if line.startswith("%T\t"):
            cur = line[3:].strip()
            tables[cur], fields = [], None
        elif line.startswith("%F\t"):
            fields = line[3:].rstrip("\r").split("\t")
        elif line.startswith("%R\t") and cur and fields:
            tables[cur].append(dict(zip(fields, line[3:].rstrip("\r").split("\t"))))
    tasks = tables.get("TASK", [])
    if not tasks:
        sys.exit(f"{path}: no TASK table found in XER")
    projects = {p.get("proj_id"): p.get("proj_short_name", p.get("proj_id")) for p in tables.get("PROJECT", [])}
    by_proj = Counter(t.get("proj_id") for t in tasks)
    if project:
        pid = next((k for k, v in projects.items() if v == project or k == project), None)
        if pid is None:
            sys.exit(f"Project '{project}' not in XER. Projects: {', '.join(projects.values())}")
    else:
        pid = by_proj.most_common(1)[0][0]
        if len(by_proj) > 1:
            print(f"NOTE: XER holds {len(by_proj)} projects; using {projects.get(pid, pid)} "
                  f"({by_proj[pid]} activities). Use --project to choose another.")
    wbs = {w["wbs_id"]: w for w in tables.get("PROJWBS", [])}

    def path_of(wid):
        names, codes, seen = [], [], set()
        while wid in wbs and wid not in seen:
            seen.add(wid)
            w = wbs[wid]
            if w.get("proj_node_flag") == "Y":
                break
            names.append(w.get("wbs_name", ""))
            codes.append(w.get("wbs_short_name", ""))
            wid = w.get("parent_wbs_id")
        return names[::-1], ".".join(codes[::-1])

    acts = []
    for t in tasks:
        if t.get("proj_id") != pid:
            continue
        names, code = path_of(t.get("wbs_id"))
        dur = to_dec(t.get("target_drtn_hr_cnt")) or ZERO
        acts.append(Activity(t.get("task_code", "").strip(), t.get("task_name", "").strip(), names, code,
                             dur, t.get("task_type", ""), t.get("status_code", "")))
    return acts, {"format": "XER", "project": projects.get(pid, pid)}


P6_ALIASES = {
    "id": ["activity id", "task_code", "act id", "activity_id", "activity code", "act_id", "id"],
    "name": ["activity name", "task_name", "activity description", "activity title", "name", "description"],
    "wbs": ["wbs", "wbs code", "wbs_id", "wbs path", "wbs id", "wbs__wbs_short_name"],
    "wbs_name": ["wbs name", "wbs__wbs_name", "wbs_name", "wbs description", "wbs path name"],
    "duration": ["original duration", "orig dur", "od", "target_drtn_hr_cnt", "duration", "planned duration",
                 "at completion duration", "original duration"],
    "type": ["activity type", "task_type", "type"],
    "status": ["activity status", "status_code", "status"],
}


def read_p6_table(path):
    tables = read_tables(path)
    names = list(tables)
    order = [n for n in names if n.upper() == "TASK"] + [n for n in names if n.upper() != "TASK"]
    for sheet in order:
        rows = tables[sheet]
        hdr = find_header(rows, P6_ALIASES, lambda r: "id" in r and "name" in r, last_wins=())
        if not hdr:
            continue
        hi, roles = hdr
        start = hi + 1
        # P6 Excel exports carry field codes on row 1 and labels on row 2
        if hi + 1 < len(rows) and any(norm_header(v) in ("activity id", "activity name") for v in rows[hi + 1]):
            start = hi + 2
        acts = []
        for row in rows[start:]:
            aid = cell_str(get(row, roles["id"]))
            name = cell_str(get(row, roles["name"]))
            if not aid or not name:
                continue
            wbs_code = cell_str(get(row, roles.get("wbs")))
            wbs_name = cell_str(get(row, roles.get("wbs_name")))
            path_names = [s.strip() for s in re.split(r"\s*(?:>|›|\\|/)\s*", wbs_name) if s.strip()] if wbs_name else []
            if not path_names and wbs_code:
                path_names = [s for s in re.split(r"[.>/\\]", wbs_code) if s]
                path_names = path_names[1:] if len(path_names) > 1 else path_names
            dur = to_dec(get(row, roles.get("duration"))) or ZERO
            acts.append(Activity(aid, name, path_names, wbs_code, dur, cell_str(get(row, roles.get("type"))),
                                 cell_str(get(row, roles.get("status")))))
        if acts:
            return acts, {"format": "P6 export table", "sheet": sheet}
    sys.exit(f"{path}: could not find an activity table (need Activity ID and Activity Name columns)")


def load_schedule(path, project=None):
    if Path(path).suffix.lower() == ".xer":
        return read_xer(path, project)
    return read_p6_table(path)


def prepare_activities(acts, lex):
    """Parse floors/zones, derive the location-free work type, and group activities."""
    seen = set()
    dupes = []
    for a in acts:
        if a.id in seen:
            dupes.append(a.id)
        seen.add(a.id)
        fl, zn, spans, _ = parse_location(a.name, trusted=True)
        worktype = strip_spans(a.name, spans)
        wbs_nonloc = []
        wf, wz = set(), set()
        for seg in a.wbs_path:
            f2, z2, sp2, _ = parse_location(seg, trusted=True)
            wf |= f2
            wz |= z2
            rest = norm_space(re.sub(r"^[\s\-–:|/,.()]+|[\s\-–:|/,.()]+$", "", strip_spans(seg, sp2)))
            if rest and re.search(r"[A-Za-z]{3,}", rest):
                wbs_nonloc.append(rest)
        a.floors = fl or wf
        a.zones = zn or wz
        wt = norm_space(re.sub(r"[\s\-–—:|/,.()_#]+", " ", worktype)).strip().lower()
        # "Lift 1", "Lift 2", "Crane A" are one work type spread over several activities
        wt = norm_space(re.sub(r"(?<![a-z0-9])(?:no\s*)?(?:\d+|[a-z])(?![a-z0-9])", " ", wt))
        if not re.search(r"[a-z]{3,}", wt):
            wt = wbs_nonloc[-1].lower() if wbs_nonloc else a.name.lower()
        a.worktype = wt
        a.wbs_nonloc = " > ".join(wbs_nonloc)
        a.group = wt
    return dupes


class Group:
    def __init__(self, key):
        self.key = key
        self.label = key
        self.acts = []
        self.vec = Counter()
        self.norm = 0.0
        self.floors = set()


def build_groups(acts, lex):
    groups = {}
    for a in acts:
        if a.kind in ("MILESTONE", "WBS_SUMMARY"):
            continue
        g = groups.setdefault(a.group, Group(a.group))
        g.acts.append(a)
        g.floors |= a.floors
    for g in groups.values():
        names = Counter(norm_space(strip_spans(a.name, parse_location(a.name, trusted=True)[2]).strip(" -–:|"))
                        for a in g.acts)
        g.label = names.most_common(1)[0][0] or g.key
        if len(names) > 1:
            g.label += f" (+{len(names) - 1} similar)"
        wbs_tok = Counter()
        for a in g.acts:
            for t in set(tokens(a.wbs_nonloc, lex)):
                wbs_tok[t] += 1
        add_tokens(g.vec, tokens(g.key, lex), 1.0)
        n = len(g.acts)
        for t, c in wbs_tok.items():
            g.vec[t] += 0.4 * (c / n) * (CANON_WEIGHT if t.startswith("_") else 1.0)
    # IDF over groups
    df = Counter()
    for g in groups.values():
        for t in g.vec:
            df[t] += 1
    N = max(len(groups), 1)
    idf = {t: math.log((N + 1) / (c + 1)) + 1 for t, c in df.items()}
    for g in groups.values():
        for t in list(g.vec):
            g.vec[t] *= idf[t]
        g.norm = math.sqrt(sum(v * v for v in g.vec.values())) or 1.0
    return groups, idf


# --------------------------------------------------------------------------
# BOQ reading and validation
# --------------------------------------------------------------------------

BOQ_ALIASES = {
    "item": ["item", "item no", "itemno", "item ref", "ref", "reference", "no", "s no", "s/no", "sr", "sr no",
             "sl no", "code", "item code", "boq ref", "boq item", "boq no", "pay item"],
    "description": ["description", "desc", "particulars", "description of work", "description of works",
                    "work description", "item description", "details", "scope", "description of item"],
    "unit": ["unit", "units", "uom", "u/m", "unit of measure", "unit of measurement"],
    "qty": ["qty", "quantity", "quantities", "quant", "qnty", "qty."],
    "rate": ["rate", "unit rate", "unit price", "price", "rate per unit", "unit cost"],
    "amount": ["amount", "total", "total amount", "value", "total price", "extension", "total cost", "amount total",
               "line total", "cost", "sum"],
    "location": ["location", "floor", "level", "zone", "area", "floor/zone", "level/zone"],
    "section": ["bill", "section", "trade", "element", "division", "category", "package", "bill no", "work package"],
}
TOTAL_RE = re.compile(r"\b(sub-?\s?total|total|carried\s+(?:forward|to)|brought\s+forward|c/f|b/f|to\s+collection|"
                      r"collection|summary|page\s+total)\b|^(?:page|pg\.?)\s*\d+\s*$", re.I)
BF_RE = re.compile(r"brought\s+forward|\bb/f\b", re.I)
FINAL_TOTAL_RE = re.compile(r"to\s+(?:the\s+)?(?:summary|collection|main\s+summary)|bill\s+total|grand\s+total|"
                            r"total\s+carried\s+to|end\s+of\s+bill|total\s+for\s+bill", re.I)
SKIP_SHEET_RE = re.compile(r"summary|collection|grand\s*total|cover|preamble|notes|instructions|index|contents", re.I)
PRELIM_RE = re.compile(r"prelim|general\s+requirements|general\s+items|provisional\s+sum|\bp\.?s\.?\b|prime\s+cost|"
                       r"\bp\.?c\.?\s+sum|contingenc|daywork|insurance|performance\s+bond|advance\s+payment\s+bond|"
                       r"overheads?\b|attendance|site\s+supervision|staff\s+cost|builder'?s'?\s+work|\bbwic\b|"
                       r"\bsundr(?:y|ies)\b", re.I)
LS_UNITS = {"item", "ls", "l.s", "l.s.", "lump sum", "lumpsum", "sum", "ps", "p.s", "prov sum", "lot", "job", "set"}
INCL_RE = re.compile(r"^(incl|included|inc|in\s+rate|rate\s+only|n/?a|excluded|by\s+others|nil|-+)\.?$", re.I)


class BoqLine:
    def __init__(self, **kw):
        self.__dict__.update(kw)
        self.issues = []
        self.status = ""
        self.reason = ""
        self.score = 0.0
        self.candidates = []
        self.group = None
        self.allocs = []
        self.floors = set()
        self.zones = set()
        self.unknown_floors = set()
        self.suggest = None
        self.prelim = False
        self.desc_tokens = set()
        self.after_total = False

    @property
    def context(self):
        parts = [self.section] + self.headings
        return " > ".join(p for p in parts if p)


class Issue:
    def __init__(self, severity, code, line, detail, sheet="", row=""):
        self.severity, self.code, self.line, self.detail = severity, code, line, detail
        self.sheet = sheet or (line.sheet if line else "")
        self.row = row or (line.row if line else "")


def read_boq(paths, args):
    lines, issues, meta = [], [], []
    multi = len(paths) > 1
    only = {s.strip().lower() for s in args.boq_sheets.split(",")} if args.boq_sheets else None
    for path in paths:
        tables = read_tables(path)
        for sheet, rows in tables.items():
            if only is not None and sheet.lower() not in only:
                continue
            if only is None and SKIP_SHEET_RE.search(sheet) and len(tables) > 1:
                meta.append({"file": Path(path).name, "sheet": sheet, "skipped": "summary/cover sheet name"})
                continue
            if args.boq_header_row:
                hi = args.boq_header_row - 1
                roles = {}
                for j, v in enumerate(rows[hi] if hi < len(rows) else []):
                    r = match_role(norm_header(v), BOQ_ALIASES)
                    if r and (r not in roles or r == "amount"):
                        roles[r] = j
            else:
                hdr = find_header(rows, BOQ_ALIASES,
                                  lambda r: "description" in r and ("amount" in r or ("qty" in r and "rate" in r)))
                if not hdr:
                    meta.append({"file": Path(path).name, "sheet": sheet, "skipped": "no BOQ header found"})
                    continue
                hi, roles = hdr
            if args.boq_cols:
                for part in args.boq_cols.split(","):
                    k, _, v = part.partition("=")
                    roles[k.strip().lower()] = col_ref_to_index(v, rows[hi] if hi < len(rows) else ())
            if "description" not in roles:
                meta.append({"file": Path(path).name, "sheet": sheet, "skipped": "no description column"})
                continue
            prefix = f"{Path(path).stem}:" if multi else ""
            info = {"file": Path(path).name, "sheet": sheet, "header_row": hi + 1,
                    "columns": {k: get_column_letter(v + 1) for k, v in sorted(roles.items(), key=lambda x: x[1])},
                    "priced": 0, "headings": 0, "totals": 0}
            parse_sheet(rows, hi, roles, sheet, prefix, lines, issues, info, args)
            meta.append(info)
    return lines, issues, meta


def parse_sheet(rows, hi, roles, sheet, prefix, lines, issues, info, args):
    tol = Decimal(str(args.tolerance))
    section, minor = sheet, []
    seg_sum = cum_sum = ZERO
    after_final = None  # set once a "Total carried to Summary"-type row is passed
    info["printed_final_total"] = None
    for ri in range(hi + 1, len(rows)):
        row = rows[ri]
        excel_row = ri + 1
        desc = norm_space(cell_str(get(row, roles.get("description"))))
        item = cell_str(get(row, roles.get("item")))
        unit = cell_str(get(row, roles.get("unit")))
        raw_q, raw_r, raw_a = get(row, roles.get("qty")), get(row, roles.get("rate")), get(row, roles.get("amount"))
        qty, rate, amt = to_dec(raw_q), to_dec(raw_r), to_dec(raw_a)
        loc = cell_str(get(row, roles.get("location")))
        sec_col = cell_str(get(row, roles.get("section")))
        if not desc and not item and qty is None and rate is None and amt is None:
            continue
        txt_amt = cell_str(raw_a) if amt is None else ""
        txt_rate = cell_str(raw_r) if rate is None else ""

        is_total = bool(desc) and TOTAL_RE.search(desc) and qty is None and rate is None
        if is_total:
            info["totals"] += 1
            if amt is not None and not BF_RE.search(desc):
                if abs(amt - seg_sum) <= tol or abs(amt - cum_sum) <= tol:
                    pass
                else:
                    issues.append(Issue("ERROR", "SUBTOTAL_MISMATCH", None,
                                        f"'{desc[:60]}' shows {money(amt)} but priced lines above sum to "
                                        f"{money(seg_sum)} (since previous total) / {money(cum_sum)} (sheet to date)",
                                        sheet, excel_row))
                seg_sum = ZERO
                if FINAL_TOTAL_RE.search(desc):
                    after_final = desc[:50]
                    info["printed_final_total"] = amt
                    section, minor = sheet, []  # don't carry the last heading past the bill total
            continue

        has_number = qty is not None or rate is not None or amt is not None
        if not has_number:
            if desc and (INCL_RE.match(txt_amt) or INCL_RE.match(txt_rate)) and txt_amt + txt_rate:
                issues.append(Issue("INFO", "INCLUDED_ELSEWHERE", None,
                                    f"'{desc[:60]}' priced as '{txt_amt or txt_rate}' - nothing to load", sheet, excel_row))
                continue
            if desc:  # heading / spec text
                info["headings"] += 1
                letters = re.sub(r"[^A-Za-z]", "", desc)
                if len(letters) >= 4 and letters.isupper() and len(desc) < 80:
                    section, minor = desc, []
                else:
                    minor = (minor + [desc[:120]])[-2:]
            continue

        if amt is None and rate is None and (INCL_RE.match(txt_amt) or INCL_RE.match(txt_rate)) and txt_amt + txt_rate:
            issues.append(Issue("INFO", "INCLUDED_ELSEWHERE", None,
                                f"'{desc[:60]}' priced as '{txt_amt or txt_rate}' - nothing to load", sheet, excel_row))
            continue

        # Priced (or meant-to-be-priced) line
        computed = cents(qty * rate) if qty is not None and rate is not None else None
        line = BoqLine(key=f"{prefix}{sheet}!R{excel_row}", sheet=sheet, row=excel_row, item=item,
                       description=desc or "(no description)", unit=unit, qty=qty, rate=rate, amount=amt,
                       section=sec_col or section, headings=list(minor), location_text=loc, value=ZERO,
                       value_source="")
        if amt is not None:
            line.value = cents(amt)
            line.value_source = "stated amount"
            if args.amount_basis == "computed" and computed is not None:
                line.value, line.value_source = computed, "qty x rate"
        elif computed is not None:
            line.value, line.value_source = computed, "qty x rate (amount blank)"
            line.issues.append(Issue("WARNING", "AMOUNT_MISSING", line,
                                     f"Amount blank - loaded qty x rate = {money(computed)}"))
        else:
            line.value_source = "unpriced"
            if qty is not None and qty != 0:
                line.issues.append(Issue("WARNING", "UNPRICED", line,
                                         f"Quantity {qty} {unit} has no rate/amount - nothing loaded"))
        if computed is not None and amt is not None and abs(computed - cents(amt)) > tol:
            line.issues.append(Issue("ERROR", "ARITHMETIC", line,
                                     f"Qty x Rate = {money(computed)} but Amount = {money(cents(amt))} "
                                     f"(diff {money(cents(amt) - computed)}); loaded {line.value_source}"))
        if line.value < 0:
            line.issues.append(Issue("WARNING", "NEGATIVE", line, f"Negative value {money(line.value)}"))
        if qty is not None and qty < 0 and line.value >= 0:
            line.issues.append(Issue("WARNING", "NEGATIVE_QTY", line, f"Negative quantity {qty}"))
        if amt is not None and amt != 0 and qty is None and unit.lower().strip(". ") not in LS_UNITS:
            line.issues.append(Issue("INFO", "AMOUNT_WITHOUT_QTY", line, "Amount with no quantity (lump sum?)"))
        if qty == 0 and amt not in (None, ZERO):
            line.issues.append(Issue("WARNING", "ZERO_QTY_WITH_AMOUNT", line, f"Qty is 0 but amount is {money(amt)}"))
        if after_final and line.value:
            line.after_total = True
            line.issues.append(Issue("WARNING", "BELOW_BILL_TOTAL", line,
                                     f"Priced row sits below '{after_final}' so it is not in the bill's carried "
                                     f"total - stray row, late addition or omission? Held for review"))
        seg_sum += line.value
        cum_sum += line.value
        info["priced"] += 1
        info["parsed_total"] = info.get("parsed_total", ZERO) + line.value
        lines.append(line)


def cross_line_checks(lines):
    issues = []
    seen = defaultdict(list)
    for ln in lines:
        if ln.value:
            seen[(ln.description.lower(), ln.unit.lower(), ln.qty, ln.rate, ln.value, ln.context.lower())].append(ln)
    for group in seen.values():
        if len(group) > 1:
            for ln in group[1:]:
                ln.issues.append(Issue("WARNING", "POSSIBLE_DUPLICATE", ln,
                                       f"Identical to {group[0].key} (same description, context, qty, rate, amount)"))
    # rate outliers: same description + unit priced at very different rates
    by_desc = defaultdict(list)
    for ln in lines:
        if ln.rate and ln.rate > 0 and len(ln.description) > 8:
            by_desc[(re.sub(r"\W+", " ", ln.description.lower()).strip(), ln.unit.lower())].append(ln)
    for group in by_desc.values():
        if len(group) < 3:
            continue
        rates = sorted(ln.rate for ln in group)
        med = rates[len(rates) // 2]
        for ln in group:
            if med > 0 and abs(ln.rate - med) / med > Decimal("0.25"):
                ln.issues.append(Issue("WARNING", "RATE_OUTLIER", ln,
                                       f"Rate {ln.rate} vs median {med} for {len(group)} identical items"))
    return issues


# --------------------------------------------------------------------------
# Matching, rules, decisions, splitting
# --------------------------------------------------------------------------

def line_vector(ln, lex, universe):
    """Token vector for a BOQ line. Location comes from the most specific text that has one:
    description / location column first, then the headings above it, then the section."""
    texts = [(ln.description, 1.0, 0), (ln.location_text, 0.3, 0), (" ".join(ln.headings), 0.7, 1), (ln.section, 0.5, 2)]
    vec = Counter()
    found = {}
    for txt, w, tier in texts:
        if not txt:
            continue
        f, z, spans, unk = parse_location(txt, universe)
        slot = found.setdefault(tier, [set(), set(), set()])
        slot[0] |= f
        slot[1] |= z
        slot[2] |= unk
        toks = tokens(strip_spans(txt, spans), lex)
        if txt is ln.description:
            ln.desc_tokens = set(toks)
            add_tokens(vec, toks, w, decay=0.08)
        else:
            add_tokens(vec, toks, w)
    floors = zones = unknown = set()
    for tier in (0, 1, 2):
        if tier in found and found[tier][0] and not floors:
            floors, unknown = found[tier][0], found[tier][2]
        if tier in found and found[tier][1] and not zones:
            zones = found[tier][1]
    if not floors:
        unknown = found.get(0, [set(), set(), set()])[2]
        # "basement car park" with no number means every basement level
        if universe and re.search(r"\bbasements?\b(?!\s*-?\d)", f"{ln.description} {ln.location_text}", re.I):
            floors = {f for f in universe if f.startswith("B") and f[1:].isdigit()}
    ln.floors, ln.zones, ln.unknown_floors = floors, zones, unknown
    return vec


def score_groups(vec, groups, idf, line_floors):
    v = {t: w * idf[t] for t, w in vec.items() if t in idf}
    vn = math.sqrt(sum(x * x for x in v.values()))
    if not vn:
        return []
    out = []
    for g in groups.values():
        dot = sum(w * g.vec.get(t, 0.0) for t, w in v.items())
        if dot <= 0:
            continue
        s = dot / (vn * g.norm)
        if line_floors and g.floors and not (line_floors & g.floors):
            s *= 0.6
        out.append((s, g))
    out.sort(key=lambda x: -x[0])
    return out


def filter_location(acts, floors, zones):
    """Activities of a work type that sit on the line's floors/zones. Returns (acts, problem)."""
    out = acts
    if floors and any(a.floors for a in out):  # work types without floors (e.g. Backfilling) take any line
        hit = [a for a in out if a.floors & floors]
        if not hit:
            return acts, "BOQ floor(s) " + summarize_floors(floors) + " not found for this work type"
        out = hit
    if zones and any(a.zones for a in out):
        hit = [a for a in out if a.zones & zones]
        if not hit:
            return out, "BOQ zone(s) " + ",".join(sorted(zones)) + " not found for this work type"
        out = hit
    return out, ""


def load_weights(path):
    rows = list(read_tables(path).values())[0]
    hdr = [norm_header(v) for v in rows[0]]
    fi = next((i for i, h in enumerate(hdr) if h in ("floor", "level")), 0)
    zi = next((i for i, h in enumerate(hdr) if h in ("zone", "area")), None)
    wi = next((i for i, h in enumerate(hdr) if h in ("weight", "area m2", "area", "gfa", "factor", "value")
               and i != zi), len(hdr) - 1)
    w = {}
    for r in rows[1:]:
        fl_txt = cell_str(get(r, fi))
        if not fl_txt:
            continue
        f, _, _, _ = parse_location(fl_txt, trusted=True)
        if not f and fl_txt.upper() in ("GF", "RF", "LG", "UG", "MZ"):
            f = {fl_txt.upper()}
        z = cell_str(get(r, zi)).upper().replace("ZONE", "").strip() if zi is not None else ""
        val = to_dec(get(r, wi))
        for fl in f:
            if val is not None:
                w[(fl, z)] = val
    return w


def activity_weight(a, basis, weights, missing):
    if basis == "duration":
        return a.duration if a.duration and a.duration > 0 else Decimal(1)
    if basis == "weights":
        total = ZERO
        found = False
        for f in a.floors or {""}:
            for z in a.zones or {""}:
                if (f, z) in weights:
                    total += weights[(f, z)]
                    found = True
                elif (f, "") in weights:
                    total += weights[(f, "")] / max(len(a.zones), 1)
                    found = True
        if not found:
            missing.add(a.id)
            return Decimal(1)
        return total
    return Decimal(1)


def load_rules(path):
    rules = []
    rows = list(read_tables(path).values())[0]
    hdr = [norm_header(v) for v in rows[0]]
    pi = hdr.index("pattern") if "pattern" in hdr else 0
    ti = hdr.index("target") if "target" in hdr else 1
    ni = hdr.index("note") if "note" in hdr else None
    for r in rows[1:]:
        pat, tgt = cell_str(get(r, pi)), cell_str(get(r, ti))
        if pat and tgt:
            rules.append((re.compile(pat, re.I), tgt, cell_str(get(r, ni)) if ni is not None else ""))
    return rules


def load_decision_file(path):
    """Read decisions/suggestions from a CSV or from the 'Review Queue' sheet of the review workbook."""
    tables = read_tables(path)
    sheet = next((s for s in tables if s.lower() == "review queue"), next(iter(tables)))
    rows = tables[sheet]
    hi = next((i for i, r in enumerate(rows[:10]) if "boq key" in [norm_header(v) for v in r]), None)
    if hi is None:
        sys.exit(f"{path}: needs a 'BOQ Key' column")
    hdr = [norm_header(v) for v in rows[hi]]

    def col(*names):
        return next((hdr.index(n) for n in names if n in hdr), None)
    ki, di = col("boq key"), col("decision")
    ai, ni = col("activity ids", "activity id", "activities"), col("planner note", "note", "reason")
    out = {}
    for r in rows[hi + 1:]:
        key, dec = cell_str(get(r, ki)), cell_str(get(r, di)).upper().replace(" ", "")
        if key and dec:
            out[key] = {"decision": dec, "ids": cell_str(get(r, ai)) if ai is not None else "",
                        "note": cell_str(get(r, ni)) if ni is not None else ""}
    return out


def resolve_ids(spec, act_by_id):
    """'A100; A2*:60' -> [(activity, weight-or-None)], unknown tokens."""
    out, unknown = [], []
    for tok in re.split(r"[;,\n]+", spec):
        tok = tok.strip()
        if not tok:
            continue
        if tok.lower().startswith("worktype:"):  # every activity of a trade, e.g. worktype:first fix
            rx = re.compile(tok[9:].strip(), re.I)
            hits = [a for a in act_by_id.values() if rx.search(a.worktype) and a.kind not in ("MILESTONE", "WBS_SUMMARY")]
            if not hits:
                unknown.append(tok)
            out.extend((a, None) for a in hits)
            continue
        aid, _, w = tok.partition(":")
        aid = aid.strip()
        wt = to_dec(w.strip().rstrip("%")) if w else None
        if any(c in aid for c in "*?["):
            # wildcards never pick milestones/WBS summaries: P6 cannot carry assignments on them
            hits = [a for k, a in act_by_id.items() if fnmatch.fnmatchcase(k.upper(), aid.upper())
                    and a.kind not in ("MILESTONE", "WBS_SUMMARY")]
            if not hits:
                unknown.append(aid)
            out.extend((a, wt) for a in hits)
        elif aid in act_by_id and act_by_id[aid].kind in ("MILESTONE", "WBS_SUMMARY"):
            unknown.append(f"{aid} (milestone - cannot hold cost)")
        elif aid in act_by_id:
            out.append((act_by_id[aid], wt))
        else:
            unknown.append(aid)
    return out, unknown


def assign(ln, acts_w, basis, weights, missing_w, status, reason):
    acts = [a for a, _ in acts_w]
    if any(w is not None for _, w in acts_w):
        ws = [w if w is not None else ZERO for _, w in acts_w]
    else:
        ws = [activity_weight(a, basis, weights, missing_w) for a in acts]
    ln.allocs = list(zip(acts, allocate(ln.value, ws)))
    ln.status, ln.reason = status, reason


def run_matching(lines, acts, groups, idf, lex, args, rules, decisions, suggestions):
    act_by_id = {a.id: a for a in acts}
    loadable = [a for a in acts if a.kind not in ("MILESTONE", "WBS_SUMMARY")]
    loe = [a for a in acts if a.kind == "LOE"]
    universe = set().union(*(a.floors for a in acts)) if acts else set()
    weights = load_weights(args.weights) if args.weights else {}
    missing_w = set()
    rule_hits = Counter()
    bad_decisions = []

    for ln in lines:
        vec = line_vector(ln, lex, universe)
        scored = score_groups(vec, groups, idf, ln.floors)
        ln.candidates = scored[:3]
        full_text = f"{ln.description} {ln.context}"
        ln.prelim = bool(PRELIM_RE.search(full_text))
        if ln.key in suggestions:
            ln.suggest = suggestions[ln.key]

        if ln.value == 0:
            ln.status, ln.reason = "NIL", "zero value"
            continue

        # 1. planner decisions
        d = decisions.get(ln.key)
        if d:
            dec = d["decision"]
            note = f"Planner: {dec}" + (f" - {d['note']}" if d["note"] else "")
            if dec == "EXCLUDE":
                ln.status, ln.reason = "EXCLUDED", note
                continue
            if dec in ("ACCEPT", "ACCEPT1", "ACCEPT2", "ACCEPT3"):
                k = int(dec[-1]) - 1 if dec[-1].isdigit() else 0
                if k < len(ln.candidates):
                    g = ln.candidates[k][1]
                    sel, _ = filter_location(g.acts, ln.floors, ln.zones)
                    ln.group = g
                    assign(ln, [(a, None) for a in sel], args.split, weights, missing_w, "DECIDED", note)
                    continue
                bad_decisions.append((ln.key, f"{dec}: no candidate #{k + 1}"))
            elif dec == "MANUAL":
                acts_w, unknown = resolve_ids(d["ids"], act_by_id)
                if unknown:
                    bad_decisions.append((ln.key, "unknown activity IDs: " + ", ".join(unknown)))
                if acts_w:
                    assign(ln, acts_w, args.split, weights, missing_w, "DECIDED", note)
                    continue
            else:
                bad_decisions.append((ln.key, f"unknown decision '{dec}' (use ACCEPT, ACCEPT2, ACCEPT3, MANUAL, EXCLUDE)"))

        # 2. rules
        matched_rule = False
        for i, (rx, target, rnote) in enumerate(rules):
            if not rx.search(full_text):
                continue
            kind, _, spec = target.partition(":")
            kind = kind.strip().lower()
            reason = f"Rule {i + 1}: /{rx.pattern}/ -> {target}" + (f" ({rnote})" if rnote else "")
            if kind == "exclude":
                ln.status, ln.reason = "EXCLUDED", reason
            elif kind == "activity":
                acts_w, unknown = resolve_ids(spec, act_by_id)
                if not acts_w:
                    continue
                assign(ln, acts_w, args.split, weights, missing_w, "RULE", reason)
            elif kind == "worktype":
                wrx = re.compile(spec.strip(), re.I)
                pool = [a for a in loadable if wrx.search(a.worktype) or wrx.search(groups[a.group].label if a.group in groups else "")]
                if not pool:
                    continue
                sel, prob = filter_location(pool, ln.floors, ln.zones)
                if prob:
                    ln.status, ln.reason = "REVIEW", f"{reason}, but {prob}"
                else:
                    assign(ln, [(a, None) for a in sel], args.split, weights, missing_w, "RULE", reason)
            else:
                continue
            rule_hits[i] += 1
            matched_rule = True
            break
        if matched_rule:
            continue

        # 3. automatic matching
        if args.prelims_activity and ln.prelim:
            acts_w, _ = resolve_ids(args.prelims_activity, act_by_id)
            if acts_w:
                assign(ln, acts_w, args.split, weights, missing_w, "RULE", "Preliminaries -> --prelims-activity")
                continue
        if ln.after_total:
            ln.status, ln.reason = "REVIEW", "Below the bill's final total (not in the contract sum) - confirm with the QS"
            continue
        if ln.prelim:
            hint = f" LOE candidates: {', '.join(a.id for a in loe[:5])}" if loe else ""
            ln.status, ln.reason = "REVIEW", "Preliminaries / provisional / non-activity cost - decide LOE, spread or exclude." + hint
            continue
        if ln.value < 0:
            ln.status, ln.reason = "REVIEW", "Negative amount (omission/adjustment) - confirm where it belongs"
            continue
        if not scored or scored[0][0] < args.min_score:
            ln.status, ln.reason = "REVIEW", "No credible match in the schedule (missing activity or unclear BOQ text)"
            continue
        s1, g1 = scored[0]
        s2 = scored[1][0] if len(scored) > 1 else 0.0
        ln.score, ln.group = s1, g1
        sel, prob = filter_location(g1.acts, ln.floors, ln.zones)
        if ln.unknown_floors:
            prob = (prob + "; " if prob else "") + "BOQ names floor(s) " + ", ".join(sorted(ln.unknown_floors)) + " that are not in the schedule"
        if prob:
            ln.status, ln.reason = "REVIEW", prob
        elif not (ln.desc_tokens & set(g1.vec)):
            ln.status, ln.reason = "REVIEW", "Matched on the bill/section heading only - the description itself names no scheduled work"
        elif s1 >= args.auto_threshold and (s1 - s2) >= args.margin:
            assign(ln, [(a, None) for a in sel], args.split, weights, missing_w, "AUTO", "")
        else:
            why = "low confidence" if s1 < args.auto_threshold else f"ambiguous: top two work types within {s1 - s2:.2f}"
            ln.status, ln.reason = "REVIEW", why

    return {"rule_hits": rule_hits, "bad_decisions": bad_decisions, "missing_weights": sorted(missing_w)}


# --------------------------------------------------------------------------
# Reconciliation and outputs
# --------------------------------------------------------------------------

def reconcile(lines, acts):
    by_status = defaultdict(lambda: [0, ZERO])
    for ln in lines:
        by_status[ln.status][0] += 1
        by_status[ln.status][1] += ln.value
    boq_total = sum((ln.value for ln in lines), ZERO)
    loaded = sum((by_status[s][1] for s in LOADED), ZERO)
    pending = by_status["REVIEW"][1]
    excluded = by_status["EXCLUDED"][1]
    alloc_total = ZERO
    bad = []
    for ln in lines:
        s = sum((v for _, v in ln.allocs), ZERO)
        alloc_total += s
        if ln.status in LOADED and s != ln.value:
            bad.append(ln.key)
    act_tot = defaultdict(lambda: ZERO)
    act_lines = defaultdict(int)
    for ln in lines:
        for a, v in ln.allocs:
            act_tot[a.id] += v
            act_lines[a.id] += 1
    raw_sum = sum((ln.amount for ln in lines if ln.amount is not None and ln.value_source == "stated amount"), ZERO)
    rounded_sum = sum((ln.value for ln in lines if ln.value_source == "stated amount"), ZERO)
    rec = {
        "boq_total": boq_total,
        "loaded": loaded,
        "pending_review": pending,
        "excluded": excluded,
        "difference": boq_total - loaded - pending - excluded,
        "allocated_to_activities": alloc_total,
        "allocation_vs_loaded": alloc_total - loaded,
        "line_level_breaks": bad,
        "sub_cent_rounding": cents(rounded_sum - raw_sum) if raw_sum else ZERO,
        "by_status": {k: {"lines": v[0], "value": v[1]} for k, v in by_status.items()},
    }
    return rec, act_tot, act_lines


HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF")
STATUS_FILL = {"AUTO": "E2EFDA", "RULE": "DDEBF7", "DECIDED": "D9E1F2", "REVIEW": "FFF2CC",
               "EXCLUDED": "EDEDED", "NIL": "F2F2F2"}
SEV_FILL = {"ERROR": "F8CBAD", "WARNING": "FFE699", "INFO": "EDEDED"}
MONEY_FMT = "#,##0.00"


def write_sheet(ws, headers, rows, widths=None, money_cols=(), freeze="A2", fill_col=None, fill_map=None):
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font = HDR_FILL, HDR_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append([float(v) if isinstance(v, Decimal) else v for v in r])
    mc = [headers.index(h) + 1 for h in money_cols if h in headers]
    fc = headers.index(fill_col) + 1 if fill_col in headers else None
    for row in ws.iter_rows(min_row=2):
        for ci in mc:
            row[ci - 1].number_format = MONEY_FMT
        if fc and fill_map:
            col = fill_map.get(row[fc - 1].value)
            if col:
                row[fc - 1].fill = PatternFill("solid", fgColor=col)
    for i, h in enumerate(headers, 1):
        w = (widths or {}).get(h, min(max(len(h) + 2, 10), 40))
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = freeze
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"


def write_review_workbook(path, lines, acts, issues, rec, act_tot, act_lines, groups, meta, run_info):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "P6 Cost Loading - Review Workbook"
    ws["A1"].font = Font(bold=True, size=14)
    rows = [
        ("Generated", run_info["generated"]),
        ("BOQ file(s)", ", ".join(run_info["boq_files"])),
        ("Schedule", run_info["schedule"]),
        ("Split basis", run_info["split"]),
        ("P6 resource", run_info["resource_id"]),
        ("", ""),
        ("RECONCILIATION", ""),
        ("Total priced BOQ value", rec["boq_total"]),
        ("Loaded to activities (AUTO + RULE + DECIDED)", rec["loaded"]),
        ("Pending planner review", rec["pending_review"]),
        ("Excluded by planner/rule", rec["excluded"]),
        ("Difference (must be 0.00)", rec["difference"]),
        ("Allocated to activities (sum of splits)", rec["allocated_to_activities"]),
        ("Allocation vs loaded (must be 0.00)", rec["allocation_vs_loaded"]),
        ("P6 import file total", run_info.get("import_total", "")),
        ("Sub-cent rounding of BOQ amounts", rec["sub_cent_rounding"]),
        ("", ""),
        ("COVERAGE", ""),
        ("Value loaded %", f"{run_info['pct_value_loaded']:.1f}%"),
        ("Priced BOQ lines", len(lines)),
        ("Activities in schedule (excl. milestones)", run_info["loadable_acts"]),
        ("Activities receiving cost", sum(1 for v in act_tot.values() if v)),
        ("Work types (activity groups)", len(groups)),
        ("", ""),
        ("STATUS", "Lines / Value"),
    ]
    for st in ("AUTO", "RULE", "DECIDED", "REVIEW", "EXCLUDED", "NIL"):
        v = rec["by_status"].get(st)
        if v:
            rows.append((st, f"{v['lines']} lines / {money(v['value'])}"))
    rows += [("", ""), ("BOQ ISSUES", "")]
    for (sev, code), n in sorted(Counter((i.severity, i.code) for i in issues).items()):
        rows.append((f"{sev}: {code}", n))
    for r, (k, v) in enumerate(rows, start=3):
        ws.cell(r, 1, k)
        c = ws.cell(r, 2, float(v) if isinstance(v, Decimal) else v)
        if isinstance(v, Decimal):
            c.number_format = MONEY_FMT
        if k.isupper() and k:
            ws.cell(r, 1).font = Font(bold=True)
        if "must be 0.00" in k:
            ok = v == 0
            c.fill = PatternFill("solid", fgColor="C6EFCE" if ok else "FFC7CE")
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 60

    # Review queue (editable)
    rq = wb.create_sheet("Review Queue")
    review = sorted([ln for ln in lines if ln.status == "REVIEW"], key=lambda l: -abs(l.value))
    hdrs = ["BOQ Key", "Item", "Description", "Context", "Unit", "Qty", "Value", "Reason", "Issues",
            "Cand 1", "Score 1", "Cand 1 Activities", "Cand 2", "Score 2", "Cand 3", "Score 3",
            "Suggested Decision", "Suggested IDs", "Suggestion Note", "Decision", "Activity IDs", "Planner Note"]
    data = []
    for ln in review:
        c = ln.candidates + [(None, None)] * (3 - len(ln.candidates))
        sel = ""
        if c[0][1]:
            acts_sel, _ = filter_location(c[0][1].acts, ln.floors, ln.zones)
            sel = f"{len(acts_sel)} acts {summarize_floors(set().union(*(a.floors for a in acts_sel)))}"
        sg = ln.suggest or {}
        data.append([ln.key, ln.item, ln.description, ln.context, ln.unit, float(ln.qty) if ln.qty is not None else None,
                     ln.value, ln.reason, "; ".join(i.code for i in ln.issues),
                     c[0][1].label if c[0][1] else "", round(c[0][0], 3) if c[0][0] is not None else None, sel,
                     c[1][1].label if c[1][1] else "", round(c[1][0], 3) if c[1][0] is not None else None,
                     c[2][1].label if c[2][1] else "", round(c[2][0], 3) if c[2][0] is not None else None,
                     sg.get("decision", ""), sg.get("ids", ""), sg.get("note", ""), "", "", ""])
    write_sheet(rq, hdrs, data, widths={"Description": 50, "Context": 35, "Reason": 40, "Cand 1": 30, "Cand 2": 30,
                                        "Cand 3": 30, "Suggestion Note": 40, "Planner Note": 30, "Activity IDs": 25,
                                        "Suggested IDs": 20, "BOQ Key": 18},
                money_cols=["Value"], freeze="D2")
    edit_fill = PatternFill("solid", fgColor="FFF2CC")
    for col in ("Decision", "Activity IDs", "Planner Note"):
        ci = hdrs.index(col) + 1
        rq.cell(1, ci).fill = PatternFill("solid", fgColor="C65911")
        for r in range(2, len(data) + 2):
            rq.cell(r, ci).fill = edit_fill
    if data:
        dv = DataValidation(type="list", formula1='"ACCEPT,ACCEPT2,ACCEPT3,MANUAL,EXCLUDE"', allow_blank=True)
        dcol = get_column_letter(hdrs.index("Decision") + 1)
        dv.add(f"{dcol}2:{dcol}{len(data) + 1}")
        rq.add_data_validation(dv)

    # All BOQ lines
    bl = wb.create_sheet("BOQ Lines")
    hdrs = ["BOQ Key", "Sheet", "Row", "Item", "Description", "Context", "Unit", "Qty", "Rate", "Amount",
            "Value Loaded", "Value Source", "Status", "Reason", "Score", "Work Type", "Activities", "Floors", "Zones", "Issues"]
    data = []
    for ln in lines:
        data.append([ln.key, ln.sheet, ln.row, ln.item, ln.description, ln.context, ln.unit,
                     float(ln.qty) if ln.qty is not None else None, float(ln.rate) if ln.rate is not None else None,
                     float(ln.amount) if ln.amount is not None else None, ln.value, ln.value_source, ln.status,
                     ln.reason, round(ln.score, 3) if ln.score else None, ln.group.label if ln.group else "",
                     len(ln.allocs) or None, summarize_floors(ln.floors), ",".join(sorted(ln.zones)),
                     "; ".join(f"{i.code}: {i.detail}" for i in ln.issues)])
    write_sheet(bl, hdrs, data, widths={"Description": 50, "Context": 35, "Reason": 40, "Work Type": 30, "Issues": 50},
                money_cols=["Rate", "Amount", "Value Loaded"], freeze="F2", fill_col="Status", fill_map=STATUS_FILL)

    # Allocations
    al = wb.create_sheet("Allocations")
    hdrs = ["BOQ Key", "Item", "Description", "Status", "Activity ID", "Activity Name", "Floor", "Zone", "Allocated"]
    data = [[ln.key, ln.item, ln.description, ln.status, a.id, a.name, summarize_floors(a.floors),
             ",".join(sorted(a.zones)), v] for ln in lines for a, v in ln.allocs]
    write_sheet(al, hdrs, data, widths={"Description": 45, "Activity Name": 40}, money_cols=["Allocated"])

    # Activity totals
    at = wb.create_sheet("Activity Totals")
    hdrs = ["Activity ID", "Activity Name", "WBS", "Type", "Floor", "Zone", "Work Type", "BOQ Lines", "Budgeted Cost"]
    data = [[a.id, a.name, a.wbs_code or " > ".join(a.wbs_path), a.kind, summarize_floors(a.floors),
             ",".join(sorted(a.zones)), a.worktype, act_lines.get(a.id, 0), act_tot.get(a.id, ZERO)]
            for a in acts if a.kind not in ("MILESTONE", "WBS_SUMMARY")]
    data.append(["TOTAL", "", "", "", "", "", "", sum(act_lines.values()), sum(act_tot.values(), ZERO)])
    write_sheet(at, hdrs, data, widths={"Activity Name": 45, "WBS": 30, "Work Type": 30}, money_cols=["Budgeted Cost"])

    # Issues
    iss = wb.create_sheet("BOQ Issues")
    hdrs = ["Severity", "Issue", "BOQ Key", "Sheet", "Row", "Item", "Description", "Value", "Detail"]
    sev_rank = {"ERROR": 0, "WARNING": 1, "INFO": 2}
    data = [[i.severity, i.code, i.line.key if i.line else "", i.sheet, i.row, i.line.item if i.line else "",
             i.line.description if i.line else "", i.line.value if i.line else None, i.detail]
            for i in sorted(issues, key=lambda i: (sev_rank[i.severity], i.code, str(i.sheet), int(i.row or 0)))]
    write_sheet(iss, hdrs, data, widths={"Description": 45, "Detail": 70}, money_cols=["Value"],
                fill_col="Severity", fill_map=SEV_FILL)

    # Activities with no cost
    un = wb.create_sheet("Unloaded Activities")
    hdrs = ["Activity ID", "Activity Name", "WBS", "Type", "Floor", "Zone", "Work Type"]
    data = [[a.id, a.name, a.wbs_code or " > ".join(a.wbs_path), a.kind, summarize_floors(a.floors),
             ",".join(sorted(a.zones)), a.worktype]
            for a in acts if a.kind not in ("MILESTONE", "WBS_SUMMARY") and not act_tot.get(a.id)]
    write_sheet(un, hdrs, data, widths={"Activity Name": 45, "WBS": 30, "Work Type": 30})

    # Parsing info
    pi = wb.create_sheet("Parsing")
    hdrs = ["File", "Sheet", "Header Row", "Columns", "Priced Lines", "Headings", "Total Rows", "Skipped"]
    data = [[m.get("file"), m.get("sheet"), m.get("header_row"), json.dumps(m.get("columns", {})),
             m.get("priced"), m.get("headings"), m.get("totals"), m.get("skipped", "")] for m in meta]
    write_sheet(pi, hdrs, data, widths={"Columns": 70})
    wb.save(path)


TASKRSRC_FIELDS = [("task_code", "Activity ID"), ("rsrc_id", "Resource ID"), ("target_qty", "Budgeted Units"),
                   ("target_cost", "Budgeted Cost")]


def write_p6_import(path, act_tot, acts, args):
    """TASKRSRC sheet: one assignment of the cost resource per activity. Returns the total written."""
    rows = [(a.id, act_tot[a.id]) for a in acts if act_tot.get(a.id)]
    if args.p6_template:
        wb = openpyxl.load_workbook(args.p6_template)
        ws = next((w for w in wb.worksheets if w.title.upper() == "TASKRSRC"), None)
        if ws is None:
            sys.exit(f"{args.p6_template}: no TASKRSRC sheet - export activities with resource assignments from P6 first")
        codes = [cell_str(c.value).lower() for c in ws[1]]
        if ws.max_row > 2:
            ws.delete_rows(3, ws.max_row - 2)
        idx = {c: i for i, c in enumerate(codes)}
        act_col = idx.get("task_code", idx.get("task_id"))
        if act_col is None or "rsrc_id" not in idx:
            sys.exit("Template TASKRSRC row 1 must contain task_code (or task_id) and rsrc_id")
        for r, (aid, v) in enumerate(rows, start=3):
            ws.cell(r, act_col + 1, aid)
            ws.cell(r, idx["rsrc_id"] + 1, args.resource_id)
            for f in ("target_qty", "target_cost", "remain_qty", "remain_cost"):
                if f in idx:
                    ws.cell(r, idx[f] + 1, v)
            if args.cost_account and "acct_id" in idx:
                ws.cell(r, idx["acct_id"] + 1, args.cost_account)
        wb.save(path)
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "TASKRSRC"
        fields = list(TASKRSRC_FIELDS) + ([("acct_id", "Cost Account")] if args.cost_account else [])
        ws.append([f for f, _ in fields])
        ws.append([l for _, l in fields])
        for aid, v in rows:
            r = [aid, args.resource_id, v, v] + ([args.cost_account] if args.cost_account else [])
            ws.append(r)
        for i in range(1, len(fields) + 1):
            ws.column_dimensions[get_column_letter(i)].width = 18
        wb.save(path)
    # read back and verify
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = next(w for w in wb.worksheets if w.title.upper() == "TASKRSRC")
    all_rows = list(ws.iter_rows(values_only=True))
    codes = [cell_str(c).lower() for c in all_rows[0]]
    ci = codes.index("target_qty") if "target_qty" in codes else codes.index("target_cost")
    total = sum((cents(to_dec(r[ci]) or ZERO) for r in all_rows[2:] if r and r[ci] is not None), ZERO)
    wb.close()
    return total, len(rows)


def write_assignments_csv(path, act_tot, act_lines, acts, args):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["Activity ID", "Activity Name", "Resource ID", "Budgeted Units", "Budgeted Cost", "BOQ Lines"])
        for a in acts:
            v = act_tot.get(a.id)
            if v:
                w.writerow([a.id, a.name, args.resource_id, f"{v:.2f}", f"{v:.2f}", act_lines[a.id]])


def auto_families(lines):
    fam = {}
    for ln in lines:
        if ln.status != "AUTO" or not ln.group:
            continue
        desc = norm_space(re.sub(r"\d+(?:\.\d+)?|zone\s+\w\b|level\s+\w+", "#", ln.description, flags=re.I))
        f = fam.setdefault((desc.lower(), ln.group.label), {"description": desc, "work_type": ln.group.label,
                                                            "lines": 0, "value": ZERO, "min_score": 1.0, "example": ln.key})
        f["lines"] += 1
        f["value"] += ln.value
        f["min_score"] = round(min(f["min_score"], ln.score), 3)
    return sorted(fam.values(), key=lambda f: f["min_score"])


def as_json(o):
    if isinstance(o, Decimal):
        return float(o)
    if isinstance(o, set):
        return sorted(o)
    raise TypeError(type(o))


# --------------------------------------------------------------------------
# Commands
# --------------------------------------------------------------------------

def load_inputs(args):
    lex = Lexicon(args.lexicon or LEXICON_PATH)
    acts, smeta = load_schedule(args.p6, args.project)
    dupes = prepare_activities(acts, lex)
    groups, idf = build_groups(acts, lex)
    lines, issues, bmeta = read_boq(args.boq, args)
    return lex, acts, smeta, dupes, groups, idf, lines, issues, bmeta


def cmd_inspect(args):
    lex, acts, smeta, dupes, groups, idf, lines, issues, bmeta = load_inputs(args)
    universe = set().union(*(a.floors for a in acts)) if acts else set()
    print("=== BOQ ===")
    for m in bmeta:
        if m.get("skipped"):
            print(f"  [{m['file']} / {m['sheet']}] SKIPPED: {m['skipped']}")
        else:
            pt = m.get("printed_final_total")
            parsed = m.get("parsed_total", ZERO)
            flag = "" if pt is None else ("  OK" if pt == parsed else f"  <-- DIFF {money(parsed - pt)}")
            print(f"  [{m['file']} / {m['sheet']}] header row {m['header_row']}  columns {m['columns']}  "
                  f"priced={m['priced']} headings={m['headings']} totals={m['totals']}")
            print(f"      parsed total {money(parsed)}   printed bill total "
                  f"{money(pt) if pt is not None else '(none found)'}{flag}")
    total = sum((ln.value for ln in lines), ZERO)
    print(f"  Priced lines: {len(lines)}   Total value: {money(total)}   "
          f"(compare with the BOQ summary / grand total page)")
    below = [ln for ln in lines if ln.after_total]
    if below:
        print(f"  WARNING {len(below)} priced rows sit below a bill total: "
              + ", ".join(f"{ln.key} {money(ln.value)}" for ln in below[:6]))
    for ln in lines[:8]:
        line_vector(ln, lex, universe)
        print(f"   - {ln.key:<16} {money(ln.value):>16}  {ln.description[:60]!r}  ctx={ln.context[:50]!r}  "
              f"floors={summarize_floors(ln.floors)} zones={','.join(sorted(ln.zones))}")
    print("\n=== SCHEDULE ===")
    kinds = Counter(a.kind for a in acts)
    print(f"  Source: {smeta}   Activities: {len(acts)}  {dict(kinds)}")
    if dupes:
        print(f"  WARNING duplicate activity IDs: {dupes[:10]}")
    fl = Counter(f for a in acts for f in a.floors)
    print(f"  Floors: {', '.join(f'{f}({fl[f]})' for f in sorted(fl, key=floor_sort_key))}")
    zn = Counter(z for a in acts for z in a.zones)
    print(f"  Zones: {', '.join(f'{z}({zn[z]})' for z in sorted(zn)) or '(none detected)'}")
    noloc = [a for a in acts if not a.floors and a.kind == "TASK"]
    print(f"  Activities without a floor: {len(noloc)}  e.g. {[a.name for a in noloc[:5]]}")
    print(f"  Work types: {len(groups)}")
    for g in sorted(groups.values(), key=lambda g: -len(g.acts))[:40]:
        print(f"   - {g.label[:50]:<50} {len(g.acts):>4} acts  floors {summarize_floors(g.floors)}")
    if len(groups) > 40:
        print(f"   ... {len(groups) - 40} more")


def cmd_run(args):
    t0 = time.time()
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    lex, acts, smeta, dupes, groups, idf, lines, issues, bmeta = load_inputs(args)
    if not lines:
        sys.exit("No priced BOQ lines found - run 'inspect' and set --boq-cols / --boq-header-row")
    issues += cross_line_checks(lines)
    rules = load_rules(args.rules) if args.rules else []
    decisions = load_decision_file(args.decisions) if args.decisions else {}
    suggestions = load_decision_file(args.suggestions) if args.suggestions else {}
    stats = run_matching(lines, acts, groups, idf, lex, args, rules, decisions, suggestions)
    for ln in lines:
        issues.extend(ln.issues)
    if dupes:
        issues.append(Issue("WARNING", "DUPLICATE_ACTIVITY_ID", None, "Duplicate activity IDs: " + ", ".join(dupes[:20]), "schedule", ""))
    for key, msg in stats["bad_decisions"]:
        issues.append(Issue("WARNING", "BAD_DECISION", None, f"{key}: {msg}", "decisions", ""))
    if stats["missing_weights"]:
        issues.append(Issue("WARNING", "MISSING_WEIGHT", None,
                            f"{len(stats['missing_weights'])} activities had no floor/zone weight (used 1): "
                            + ", ".join(stats["missing_weights"][:15]), "weights", ""))

    rec, act_tot, act_lines = reconcile(lines, acts)
    name = args.name
    import_path = out / f"{name}_P6_import_TASKRSRC.xlsx"
    import_total, n_assign = write_p6_import(import_path, act_tot, acts, args)
    write_assignments_csv(out / f"{name}_assignments.csv", act_tot, act_lines, acts, args)
    loadable = sum(1 for a in acts if a.kind not in ("MILESTONE", "WBS_SUMMARY"))
    pct_val = float(rec["loaded"] / rec["boq_total"] * 100) if rec["boq_total"] else 0.0
    run_info = {
        "generated": time.strftime("%Y-%m-%d %H:%M"),
        "boq_files": [Path(p).name for p in args.boq],
        "schedule": f"{Path(args.p6).name} ({smeta.get('format')}{', project ' + str(smeta['project']) if smeta.get('project') else ''})",
        "split": args.split + (f" ({Path(args.weights).name})" if args.weights else ""),
        "resource_id": args.resource_id,
        "import_total": import_total,
        "pct_value_loaded": pct_val,
        "loadable_acts": loadable,
    }
    review_path = out / f"{name}_review.xlsx"
    write_review_workbook(review_path, lines, acts, issues, rec, act_tot, act_lines, groups, bmeta, run_info)

    checks = {
        "boq_equals_loaded_plus_pending_plus_excluded": rec["difference"] == 0,
        "allocations_equal_loaded": rec["allocation_vs_loaded"] == 0,
        "import_file_equals_loaded": import_total == rec["loaded"],
        "no_line_level_breaks": not rec["line_level_breaks"],
    }
    elapsed = time.time() - t0
    review = sorted([ln for ln in lines if ln.status == "REVIEW"], key=lambda l: -abs(l.value))
    summary = {
        "elapsed_seconds": round(elapsed, 2),
        "files": {"review_workbook": str(review_path), "p6_import": str(import_path),
                  "assignments_csv": str(out / f"{name}_assignments.csv")},
        "counts": {"priced_lines": len(lines), "activities": len(acts), "loadable_activities": loadable,
                   "activities_costed": sum(1 for v in act_tot.values() if v), "p6_assignments": n_assign,
                   "work_types": len(groups)},
        "pct_value_loaded": round(pct_val, 2),
        "reconciliation": {k: v for k, v in rec.items() if k != "line_level_breaks"},
        "checks": checks,
        "issues": dict(Counter(f"{i.severity}:{i.code}" for i in issues)),
        "top_issues": [{"severity": i.severity, "code": i.code, "key": i.line.key if i.line else f"{i.sheet}!R{i.row}",
                        "detail": i.detail} for i in sorted(issues, key=lambda i: {"ERROR": 0, "WARNING": 1, "INFO": 2}[i.severity])[:40]],
        "rule_hits": {str(k + 1): v for k, v in stats["rule_hits"].items()},
        "bill_totals": [{"sheet": m["sheet"], "parsed": m.get("parsed_total", ZERO), "printed": m.get("printed_final_total")}
                        for m in bmeta if not m.get("skipped")],
        # one row per distinct (description, work type) among loaded lines, weakest first - for spot checks
        "auto_check": auto_families(lines),
        "review_queue": [{
            "key": ln.key, "item": ln.item, "description": ln.description[:160], "context": ln.context[:120],
            "unit": ln.unit, "qty": ln.qty, "value": ln.value, "reason": ln.reason,
            "floors": summarize_floors(ln.floors), "zones": sorted(ln.zones),
            "candidates": [{"rank": i + 1, "work_type": g.label, "score": round(s, 3), "activities": len(g.acts),
                            "floors": summarize_floors(g.floors), "example_ids": [a.id for a in g.acts[:3]]}
                           for i, (s, g) in enumerate(ln.candidates)],
        } for ln in review],
    }
    (out / f"{name}_summary.json").write_text(json.dumps(summary, default=as_json, indent=1), encoding="utf-8")

    st = rec["by_status"]
    print(f"Done in {elapsed:.1f}s")
    print(f"  Priced BOQ lines : {len(lines)}  |  Activities: {loadable} loadable, "
          f"{summary['counts']['activities_costed']} costed  |  Work types: {len(groups)}")
    for s in ("AUTO", "RULE", "DECIDED", "REVIEW", "EXCLUDED", "NIL"):
        if s in st:
            print(f"  {s:<9} {st[s]['lines']:>6} lines  {money(st[s]['value']):>20}")
    print(f"  BOQ total        {money(rec['boq_total']):>20}")
    print(f"  Loaded           {money(rec['loaded']):>20}  ({pct_val:.1f}% of value)")
    print(f"  Difference       {money(rec['difference']):>20}  (BOQ - loaded - pending - excluded)")
    print(f"  Import file total{money(import_total):>20}  ({n_assign} assignments)")
    print(f"  Checks: " + ", ".join(f"{k}={'OK' if v else 'FAIL'}" for k, v in checks.items()))
    ic = Counter(i.severity for i in issues)
    print(f"  BOQ issues: {ic.get('ERROR', 0)} errors, {ic.get('WARNING', 0)} warnings, {ic.get('INFO', 0)} info")
    if stats["bad_decisions"]:
        print(f"  WARNING: {len(stats['bad_decisions'])} decisions could not be applied (see BOQ Issues: BAD_DECISION)")
    print(f"  Outputs in {out}")
    if not all(checks.values()):
        sys.exit("RECONCILIATION FAILED - do not import; investigate the Summary sheet")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)

    def common(p):
        p.add_argument("--boq", nargs="+", required=True, help="BOQ file(s): .xlsx/.xlsm/.csv")
        p.add_argument("--p6", required=True, help="Schedule: .xer, or P6 Excel/CSV activity export")
        p.add_argument("--project", help="XER project short name when the XER holds several")
        p.add_argument("--boq-sheets", help="Comma list of BOQ sheets to read (default: all except summary/cover)")
        p.add_argument("--boq-cols", help="Override columns, e.g. 'description=C,qty=E,rate=F,amount=G,item=A,unit=D'")
        p.add_argument("--boq-header-row", type=int, help="1-based header row (default: auto-detect)")
        p.add_argument("--amount-basis", choices=["stated", "computed"], default="stated",
                       help="Load the stated Amount (default) or Qty x Rate when they disagree")
        p.add_argument("--tolerance", type=float, default=0.01, help="Arithmetic tolerance (default 0.01)")
        p.add_argument("--lexicon", help="Custom trade lexicon JSON (default assets/trade_lexicon.json)")

    common(sub.add_parser("inspect", help="Show how inputs were parsed"))
    r = sub.add_parser("run", help="Match, split, reconcile and write outputs")
    common(r)
    r.add_argument("--out", required=True)
    r.add_argument("--name", default="cost_loading", help="Output file prefix")
    r.add_argument("--split", choices=["equal", "duration", "weights"], default="equal")
    r.add_argument("--weights", help="CSV floor,zone,weight (e.g. floor areas) for --split weights")
    r.add_argument("--rules", help="CSV pattern,target,note mapping rules")
    r.add_argument("--decisions", help="Planner decisions: the edited review workbook or a CSV")
    r.add_argument("--suggestions", help="Suggested decisions (CSV) shown in the Review Queue but not applied")
    r.add_argument("--prelims-activity", help="Activity ID(s) (wildcards ok) that receive preliminaries")
    r.add_argument("--resource-id", default="BOQ-COST", help="P6 nonlabor resource carrying the cost (price/unit 1.00)")
    r.add_argument("--cost-account", help="P6 cost account ID to put on every assignment")
    r.add_argument("--p6-template", help="TASKRSRC workbook exported from your P6 to copy the exact layout from")
    r.add_argument("--auto-threshold", type=float, default=0.45, help="Min score to auto-load (default 0.45)")
    r.add_argument("--margin", type=float, default=0.08, help="Min lead over 2nd work type to auto-load (default 0.08)")
    r.add_argument("--min-score", type=float, default=0.15, help="Below this a line is 'no credible match'")
    args = ap.parse_args()
    if args.cmd == "run" and args.split == "weights" and not args.weights:
        ap.error("--split weights needs --weights FILE")
    (cmd_inspect if args.cmd == "inspect" else cmd_run)(args)


if __name__ == "__main__":
    main()
